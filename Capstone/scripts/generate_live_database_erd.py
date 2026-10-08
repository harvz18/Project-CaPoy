from __future__ import annotations

import html
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
EXPORT_DIR = ROOT / "docs" / "erd_exports"
SCHEMA_PATH = EXPORT_DIR / "live_public_schema.json"
DRAWIO_PATH = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD.drawio"
PDF_PATH = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD_MANUSCRIPT.pdf"
DOCX_PATH = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD_MANUSCRIPT.docx"
NOTES_PATH = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD_NOTES.md"
GENERATED_DATE = "2026-10-08"
PROJECT_REF = "vkjmyyrxxzznbrgxzfyn"

PAGE_W = 3508
PAGE_H = 2480

DOMAIN_TABLES = {
    "Identity and Access": [
        "profiles", "roles", "permissions", "role_permissions", "user_roles",
        "user_permissions", "provider_profiles", "provider_profile_merge_archive",
    ],
    "Marketplace and Availability": [
        "service_categories", "services", "service_packages", "service_package_items",
        "provider_availability", "provider_operating_hours",
        "provider_notification_preferences", "provider_service_listing_drafts",
        "coordinator_service_profiles", "coordinator_availability",
    ],
    "Event Planning and Coordination": [
        "events", "event_requirements", "event_budget_items",
        "event_service_selections", "event_provider_instructions",
        "event_coordinator_instructions", "event_schedule_checks",
        "event_schedule_check_results", "coordination_tasks",
        "coordinator_assignment_attempts", "coordinator_packages",
        "coordinator_package_items", "event_coordinator_package_selections",
        "event_coordinator_package_services",
    ],
    "Booking and Finance": [
        "bookings", "payments", "financial_transactions", "cash_remittances",
        "provider_payment_receipts", "provider_payout_accounts",
        "provider_payout_requests",
    ],
    "Communication, Feedback, and Operations": [
        "conversations", "conversation_participants", "messages", "notifications",
        "reviews", "coordinator_reviews", "event_feedback",
        "service_review_summaries", "support_tickets", "support_messages",
        "audit_logs", "system_settings",
    ],
}

DOMAIN_SLUGS = {
    "Identity and Access": "01_IDENTITY_ACCESS",
    "Marketplace and Availability": "02_MARKETPLACE_AVAILABILITY",
    "Event Planning and Coordination": "03_EVENT_PLANNING_COORDINATION",
    "Booking and Finance": "04_BOOKING_FINANCE",
    "Communication, Feedback, and Operations": "05_COMMUNICATION_QUALITY_OPERATIONS",
}


def load_schema() -> dict:
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    names = {table["name"] for table in schema["tables"]}
    assigned = {name for tables in DOMAIN_TABLES.values() for name in tables}
    missing = sorted(names - assigned)
    unknown = sorted(assigned - names)
    if missing or unknown:
        raise RuntimeError(f"Domain mapping mismatch. Missing={missing}; unknown={unknown}")
    return schema


def safe_id(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]+", "_", value)


def compact_type(value: str) -> str:
    value = value.replace("timestamp with time zone", "timestamptz")
    value = value.replace("timestamp without time zone", "timestamp")
    value = value.replace("character varying", "varchar")
    return value


def table_map(schema: dict) -> dict[str, dict]:
    return {table["name"]: table for table in schema["tables"]}


def fk_column_map(schema: dict) -> dict[str, set[str]]:
    result: dict[str, set[str]] = defaultdict(set)
    for fk in schema["foreign_keys"]:
        result[fk["source_table"]].update(fk["source_columns"])
    return result


def column_prefix(column: dict, is_fk: bool) -> str:
    marks = []
    if column["primary_key"]:
        marks.append("PK")
    if is_fk:
        marks.append("FK")
    if column["unique"] and not column["primary_key"]:
        marks.append("UQ")
    return "/".join(marks) if marks else ""


def drawio_column_lines(table: dict, fk_columns: set[str]) -> list[str]:
    lines = []
    for column in table["columns"]:
        prefix = column_prefix(column, column["name"] in fk_columns)
        null_mark = " ?" if column["nullable"] else ""
        label = f"{prefix + ' ' if prefix else ''}{column['name']} : {compact_type(column['type'])}{null_mark}"
        lines.append(label)
    return lines


def manuscript_column_lines(table: dict, fk_columns: set[str], limit: int = 10) -> list[str]:
    columns = table["columns"]
    keyed = [
        c for c in columns
        if c["primary_key"] or c["name"] in fk_columns
    ]
    ordinary = [
        c for c in columns
        if c not in keyed and c["name"] not in {"created_at", "updated_at"}
    ]
    timestamps = [
        c for c in columns
        if c not in keyed and c["name"] in {"created_at", "updated_at"}
    ]
    selected = []
    seen = set()
    for column in keyed + ordinary + timestamps:
        if column["name"] in seen:
            continue
        selected.append(column)
        seen.add(column["name"])
        if len(selected) >= limit:
            break
    lines = []
    for column in selected:
        prefix = column_prefix(column, column["name"] in fk_columns)
        null_mark = " ?" if column["nullable"] else ""
        lines.append(
            f"{prefix + ' ' if prefix else ''}{column['name']} : "
            f"{compact_type(column['type'])}{null_mark}"
        )
    omitted = len(columns) - len(selected)
    if omitted:
        lines.append(f"... +{omitted} additional columns")
    return lines


def multiplicities(fk: dict) -> tuple[str, str]:
    child = "0..1" if fk["source_unique"] else "0..*"
    parent = "0..1" if fk["source_nullable"] else "1"
    return child, parent


def domain_for_table(table_name: str) -> str:
    for domain, names in DOMAIN_TABLES.items():
        if table_name in names:
            return domain
    raise KeyError(table_name)


def create_drawio(schema: dict) -> None:
    tables = table_map(schema)
    fk_columns = fk_column_map(schema)
    mxfile = ET.Element(
        "mxfile",
        {
            "host": "app.diagrams.net",
            "modified": f"{GENERATED_DATE}T00:00:00.000Z",
            "agent": "Codex",
            "version": "24.7.17",
            "type": "device",
            "compressed": "false",
            "pages": str(len(DOMAIN_TABLES) + 1),
        },
    )

    add_overview_drawio_page(mxfile, schema)
    for index, (domain, names) in enumerate(DOMAIN_TABLES.items(), start=1):
        add_detail_drawio_page(
            mxfile, schema, tables, fk_columns, domain, names, index
        )

    ET.indent(mxfile, space="  ")
    xml = ET.tostring(mxfile, encoding="unicode")
    DRAWIO_PATH.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n' + xml + "\n",
        encoding="utf-8",
    )


def new_drawio_model(diagram: ET.Element, page_width: int, page_height: int):
    model = ET.SubElement(
        diagram,
        "mxGraphModel",
        {
            "dx": "1422", "dy": "794", "grid": "1", "gridSize": "10",
            "guides": "1", "tooltips": "1", "connect": "1", "arrows": "1",
            "fold": "1", "page": "1", "pageScale": "1",
            "pageWidth": str(page_width), "pageHeight": str(page_height),
            "math": "0", "shadow": "0",
        },
    )
    root = ET.SubElement(model, "root")
    ET.SubElement(root, "mxCell", {"id": "0"})
    ET.SubElement(root, "mxCell", {"id": "1", "parent": "0"})
    return model, root


def add_drawio_vertex(
    root: ET.Element,
    cell_id: str,
    value: str,
    style: str,
    x: float,
    y: float,
    width: float,
    height: float,
    parent: str = "1",
) -> ET.Element:
    cell = ET.SubElement(
        root,
        "mxCell",
        {
            "id": cell_id,
            "value": value,
            "style": style,
            "vertex": "1",
            "parent": parent,
        },
    )
    ET.SubElement(
        cell,
        "mxGeometry",
        {
            "x": str(round(x, 2)), "y": str(round(y, 2)),
            "width": str(round(width, 2)), "height": str(round(height, 2)),
            "as": "geometry",
        },
    )
    return cell


def add_overview_drawio_page(mxfile: ET.Element, schema: dict) -> None:
    diagram = ET.SubElement(
        mxfile, "diagram", {"id": "overview", "name": "1 - Manuscript Overview"}
    )
    _, root = new_drawio_model(diagram, 1800, 1200)
    add_drawio_vertex(
        root, "overview_title",
        "MULTIVENT Live Database ERD - Condensed Schema Map",
        "text;html=1;align=center;verticalAlign=middle;fontSize=24;fontStyle=1;"
        "fontColor=#000000;strokeColor=none;fillColor=none;",
        250, 25, 1300, 45,
    )
    positions = {
        "Identity and Access": (60, 120, 480, 360),
        "Marketplace and Availability": (660, 120, 480, 410),
        "Booking and Finance": (1260, 120, 480, 330),
        "Event Planning and Coordination": (360, 680, 580, 430),
        "Communication, Feedback, and Operations": (1080, 650, 580, 430),
    }
    cell_ids = {}
    for domain, names in DOMAIN_TABLES.items():
        x, y, width, height = positions[domain]
        cell_id = "overview_" + safe_id(domain)
        cell_ids[domain] = cell_id
        body = "<b>" + html.escape(domain) + "</b><br><br>" + "<br>".join(
            html.escape(name) for name in names
        )
        add_drawio_vertex(
            root, cell_id, body,
            "rounded=0;whiteSpace=wrap;html=1;align=left;verticalAlign=top;"
            "spacing=14;fontSize=13;fontColor=#000000;fillColor=#ffffff;"
            "strokeColor=#000000;strokeWidth=2;",
            x, y, width, height,
        )
    counts = Counter()
    for fk in schema["foreign_keys"]:
        source_domain = domain_for_table(fk["source_table"])
        target_domain = domain_for_table(fk["target_table"])
        if source_domain != target_domain:
            counts[(source_domain, target_domain)] += 1
    for index, ((source_domain, target_domain), count) in enumerate(
        sorted(counts.items()), start=1
    ):
        cell = ET.SubElement(
            root, "mxCell",
            {
                "id": f"overview_edge_{index}",
                "value": f"{count} FK link{'s' if count != 1 else ''}",
                "style": "edgeStyle=orthogonalEdgeStyle;rounded=0;html=1;"
                "strokeColor=#000000;strokeWidth=1.5;endArrow=classic;"
                "endFill=1;fontSize=11;labelBackgroundColor=#ffffff;",
                "edge": "1", "parent": "1",
                "source": cell_ids[source_domain],
                "target": cell_ids[target_domain],
            },
        )
        ET.SubElement(cell, "mxGeometry", {"relative": "1", "as": "geometry"})
    add_drawio_vertex(
        root, "overview_note",
        f"Live public schema | {len(schema['tables'])} tables | "
        f"{sum(len(t['columns']) for t in schema['tables'])} columns | "
        f"{len(schema['foreign_keys'])} foreign keys | Generated {GENERATED_DATE}<br>"
        "Detailed attribute-level ERDs are provided on the following pages.",
        "text;html=1;align=center;verticalAlign=middle;fontSize=12;"
        "fontColor=#000000;strokeColor=none;fillColor=none;",
        250, 1130, 1300, 45,
    )


def drawio_table_layout(names: list[str], tables: dict[str, dict]) -> tuple[dict, int, int]:
    columns = 4 if len(names) >= 10 else 3
    width = 315
    gap_x = 125
    margin_x = 70
    top = 120
    heights = [top] * columns
    positions = {}
    for name in names:
        line_count = len(tables[name]["columns"])
        height = max(120, 52 + line_count * 17)
        column = min(range(columns), key=lambda index: heights[index])
        x = margin_x + column * (width + gap_x)
        y = heights[column]
        positions[name] = (x, y, width, height)
        heights[column] += height + 90
    page_width = margin_x * 2 + columns * width + (columns - 1) * gap_x
    page_height = max(heights) + 170
    return positions, page_width, page_height


def edge_points(source_box, target_box, index: int):
    sx, sy, sw, sh = source_box
    tx, ty, tw, th = target_box
    source_y = sy + sh / 2
    target_y = ty + th / 2
    if sx + sw < tx:
        start = (sx + sw, source_y)
        end = (tx, target_y)
        middle_x = (start[0] + end[0]) / 2 + ((index % 7) - 3) * 7
    elif tx + tw < sx:
        start = (sx, source_y)
        end = (tx + tw, target_y)
        middle_x = (start[0] + end[0]) / 2 + ((index % 7) - 3) * 7
    else:
        start = (sx + sw, source_y)
        end = (tx + tw, target_y)
        middle_x = max(start[0], end[0]) + 55 + (index % 10) * 9
    return start, end, [(middle_x, start[1]), (middle_x, end[1])]


def add_detail_drawio_page(
    mxfile: ET.Element,
    schema: dict,
    tables: dict[str, dict],
    fk_columns: dict[str, set[str]],
    domain: str,
    names: list[str],
    page_index: int,
) -> None:
    positions, page_width, page_height = drawio_table_layout(names, tables)
    diagram = ET.SubElement(
        mxfile,
        "diagram",
        {"id": f"domain_{page_index}", "name": f"{page_index + 1} - {domain}"},
    )
    _, root = new_drawio_model(diagram, page_width, page_height)
    add_drawio_vertex(
        root, f"title_{page_index}",
        f"MULTIVENT Live ERD - {html.escape(domain)}",
        "text;html=1;align=center;verticalAlign=middle;fontSize=22;fontStyle=1;"
        "fontColor=#000000;strokeColor=none;fillColor=none;",
        160, 25, page_width - 320, 45,
    )
    table_ids = {name: f"table_{page_index}_{safe_id(name)}" for name in names}
    internal_fks = [
        fk for fk in schema["foreign_keys"]
        if fk["source_table"] in table_ids and fk["target_table"] in table_ids
    ]
    for edge_index, fk in enumerate(internal_fks, start=1):
        child_mult, parent_mult = multiplicities(fk)
        start_arrow = "ERzeroToOne" if fk["source_unique"] else "ERzeroToMany"
        end_arrow = "ERzeroToOne" if fk["source_nullable"] else "ERone"
        start, end, points = edge_points(
            positions[fk["source_table"]],
            positions[fk["target_table"]],
            edge_index,
        )
        label = (
            f"{child_mult} : {parent_mult} "
            f"[{', '.join(fk['source_columns'])}]"
        )
        cell = ET.SubElement(
            root, "mxCell",
            {
                "id": f"edge_{page_index}_{edge_index}",
                "value": html.escape(label),
                "style": "edgeStyle=orthogonalEdgeStyle;rounded=0;html=1;"
                "orthogonalLoop=1;jettySize=auto;strokeColor=#000000;"
                f"strokeWidth=1.4;startArrow={start_arrow};startFill=0;"
                f"endArrow={end_arrow};endFill=0;fontSize=10;"
                "labelBackgroundColor=#ffffff;",
                "edge": "1", "parent": "1",
                "source": table_ids[fk["source_table"]],
                "target": table_ids[fk["target_table"]],
            },
        )
        geometry = ET.SubElement(
            cell, "mxGeometry", {"relative": "1", "as": "geometry"}
        )
        array = ET.SubElement(geometry, "Array", {"as": "points"})
        for x, y in points:
            ET.SubElement(
                array, "mxPoint", {"x": str(round(x, 2)), "y": str(round(y, 2))}
            )
    for name in names:
        table = tables[name]
        x, y, width, height = positions[name]
        status = "RLS" if table["rls_enabled"] else "RLS OFF"
        header = f"<b>{html.escape(name)}</b> <font size=\"2\">[{status}]</font>"
        table_cell = add_drawio_vertex(
            root, table_ids[name], header,
            "swimlane;html=1;rounded=0;startSize=34;horizontal=1;"
            "fillColor=#000000;swimlaneFillColor=#ffffff;"
            "fontColor=#ffffff;strokeColor=#000000;strokeWidth=2;"
            "fontSize=14;fontStyle=1;whiteSpace=wrap;",
            x, y, width, height,
        )
        body = "<br>".join(
            html.escape(line) for line in drawio_column_lines(table, fk_columns[name])
        )
        add_drawio_vertex(
            root, table_ids[name] + "_body", body,
            "text;html=1;align=left;verticalAlign=top;spacing=7;"
            "fontSize=11;fontColor=#000000;strokeColor=none;fillColor=none;"
            "whiteSpace=wrap;overflow=hidden;",
            0, 34, width, height - 34, parent=table_ids[name],
        )
    cross_fks = [
        fk for fk in schema["foreign_keys"]
        if (fk["source_table"] in table_ids) != (fk["target_table"] in table_ids)
    ]
    note = (
        f"Legend: PK = primary key; FK = foreign key; UQ = unique; ? = nullable. "
        f"Crow's-foot labels show child : parent cardinality. "
        f"{len(cross_fks)} cross-domain foreign keys are summarized on the overview page."
    )
    add_drawio_vertex(
        root, f"note_{page_index}", html.escape(note),
        "text;html=1;align=center;verticalAlign=middle;fontSize=11;"
        "fontColor=#000000;strokeColor=none;fillColor=none;",
        80, page_height - 100, page_width - 160, 45,
    )


def load_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


FONT_TITLE = load_font(48, True)
FONT_SUBTITLE = load_font(27)
FONT_DOMAIN = load_font(31, True)
FONT_HEADER = load_font(24, True)
FONT_BODY = load_font(19)
FONT_BODY_BOLD = load_font(19, True)
FONT_SMALL = load_font(17)


def trim_text(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> str:
    if draw.textlength(text, font=font) <= width:
        return text
    candidate = text
    while candidate and draw.textlength(candidate + "...", font=font) > width:
        candidate = candidate[:-1]
    return candidate.rstrip() + "..."


def manuscript_layout(names: list[str]) -> dict[str, tuple[int, int, int, int]]:
    cols = 3
    margin_x = 100
    gap_x = 55
    box_w = (PAGE_W - 2 * margin_x - (cols - 1) * gap_x) // cols
    top = 245
    bottom = 190
    rows = math.ceil(len(names) / cols)
    gap_y = 45
    box_h = (PAGE_H - top - bottom - (rows - 1) * gap_y) // rows
    positions = {}
    for index, name in enumerate(names):
        row, col = divmod(index, cols)
        positions[name] = (
            margin_x + col * (box_w + gap_x),
            top + row * (box_h + gap_y),
            box_w,
            box_h,
        )
    return positions


def raster_route(source_box, target_box, index: int):
    sx, sy, sw, sh = source_box
    tx, ty, tw, th = target_box
    if sx + sw <= tx:
        start = (sx + sw, sy + sh // 2)
        end = (tx, ty + th // 2)
        middle_x = (start[0] + end[0]) // 2 + ((index % 5) - 2) * 8
        points = [start, (middle_x, start[1]), (middle_x, end[1]), end]
    elif tx + tw <= sx:
        start = (sx, sy + sh // 2)
        end = (tx + tw, ty + th // 2)
        middle_x = (start[0] + end[0]) // 2 + ((index % 5) - 2) * 8
        points = [start, (middle_x, start[1]), (middle_x, end[1]), end]
    else:
        start = (sx + sw // 2, sy + sh)
        end = (tx + tw // 2, ty)
        if end[1] < start[1]:
            start = (sx + sw // 2, sy)
            end = (tx + tw // 2, ty + th)
        middle_y = (start[1] + end[1]) // 2 + ((index % 5) - 2) * 6
        points = [start, (start[0], middle_y), (end[0], middle_y), end]
    return points


def label_box(draw, center, text, font=FONT_SMALL):
    bbox = draw.textbbox((0, 0), text, font=font)
    width = bbox[2] - bbox[0] + 10
    height = bbox[3] - bbox[1] + 6
    x = int(center[0] - width / 2)
    y = int(center[1] - height / 2)
    draw.rectangle((x, y, x + width, y + height), fill="white", outline="black", width=1)
    draw.text((x + 5, y + 2), text, fill="black", font=font)


def svg_escape(value: str) -> str:
    return html.escape(value, quote=True)


def generate_overview(schema: dict) -> tuple[Path, Path]:
    png_path = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD_OVERVIEW.png"
    svg_path = EXPORT_DIR / "MULTIVENT_LIVE_DATABASE_ERD_OVERVIEW.svg"
    image = Image.new("RGB", (PAGE_W, PAGE_H), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (PAGE_W // 2, 45),
        "MULTIVENT Live Database Entity-Relationship Overview",
        anchor="ma", fill="black", font=FONT_TITLE,
    )
    subtitle = (
        f"Public schema | {len(schema['tables'])} tables | "
        f"{sum(len(t['columns']) for t in schema['tables'])} columns | "
        f"{len(schema['foreign_keys'])} foreign keys"
    )
    draw.text(
        (PAGE_W // 2, 110), subtitle,
        anchor="ma", fill="black", font=FONT_SUBTITLE,
    )
    positions = {
        "Identity and Access": (80, 260, 1000, 720),
        "Marketplace and Availability": (1254, 260, 1000, 820),
        "Booking and Finance": (2428, 260, 1000, 650),
        "Event Planning and Coordination": (420, 1420, 1220, 780),
        "Communication, Feedback, and Operations": (1868, 1420, 1220, 780),
    }
    counts = Counter()
    for fk in schema["foreign_keys"]:
        source_domain = domain_for_table(fk["source_table"])
        target_domain = domain_for_table(fk["target_table"])
        if source_domain != target_domain:
            counts[(source_domain, target_domain)] += 1
    for index, ((source_domain, target_domain), count) in enumerate(
        sorted(counts.items()), start=1
    ):
        points = raster_route(positions[source_domain], positions[target_domain], index)
        draw.line(points, fill="black", width=3, joint="curve")
        segment_start = points[1]
        segment_end = points[2]
        label_box(
            draw,
            ((segment_start[0] + segment_end[0]) // 2,
             (segment_start[1] + segment_end[1]) // 2),
            f"{count} FK",
        )
    for domain, names in DOMAIN_TABLES.items():
        x, y, width, height = positions[domain]
        draw.rectangle((x, y, x + width, y + height), fill="white", outline="black", width=5)
        draw.rectangle((x, y, x + width, y + 70), fill="black", outline="black")
        draw.text(
            (x + width // 2, y + 35), domain,
            anchor="mm", fill="white", font=FONT_DOMAIN,
        )
        midpoint = math.ceil(len(names) / 2)
        columns = [names[:midpoint], names[midpoint:]]
        col_width = (width - 60) // 2
        for col_index, table_names in enumerate(columns):
            text_x = x + 30 + col_index * col_width
            for row_index, table_name in enumerate(table_names):
                draw.text(
                    (text_x, y + 105 + row_index * 46),
                    trim_text(draw, table_name, FONT_BODY, col_width - 20),
                    fill="black", font=FONT_BODY,
                )
    draw.text(
        (PAGE_W // 2, PAGE_H - 105),
        "Domain connectors show the number of cross-domain foreign keys. "
        "Attribute-level cardinalities appear in the detailed ERD pages.",
        anchor="mm", fill="black", font=FONT_SMALL,
    )
    image.save(png_path, dpi=(300, 300))

    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{PAGE_W}" height="{PAGE_H}" '
        f'viewBox="0 0 {PAGE_W} {PAGE_H}">',
        '<rect width="100%" height="100%" fill="white"/>',
        '<style>text{font-family:Arial,sans-serif;fill:#000}'
        '.title{font-size:48px;font-weight:700}.subtitle{font-size:27px}'
        '.domain{font-size:31px;font-weight:700;fill:#fff}'
        '.body{font-size:19px}.small{font-size:17px}</style>',
        f'<text class="title" x="{PAGE_W/2}" y="85" text-anchor="middle">'
        'MULTIVENT Live Database Entity-Relationship Overview</text>',
        f'<text class="subtitle" x="{PAGE_W/2}" y="135" text-anchor="middle">'
        f'{svg_escape(subtitle)}</text>',
    ]
    for index, ((source_domain, target_domain), count) in enumerate(
        sorted(counts.items()), start=1
    ):
        points = raster_route(positions[source_domain], positions[target_domain], index)
        point_text = " ".join(f"{x},{y}" for x, y in points)
        svg.append(
            f'<polyline points="{point_text}" fill="none" stroke="#000" stroke-width="3"/>'
        )
        a, b = points[1], points[2]
        lx, ly = (a[0] + b[0]) // 2, (a[1] + b[1]) // 2
        svg.append(
            f'<rect x="{lx-34}" y="{ly-15}" width="68" height="30" '
            'fill="#fff" stroke="#000"/>'
        )
        svg.append(
            f'<text class="small" x="{lx}" y="{ly+6}" text-anchor="middle">{count} FK</text>'
        )
    for domain, names in DOMAIN_TABLES.items():
        x, y, width, height = positions[domain]
        svg.append(
            f'<rect x="{x}" y="{y}" width="{width}" height="{height}" '
            'fill="#fff" stroke="#000" stroke-width="5"/>'
        )
        svg.append(
            f'<rect x="{x}" y="{y}" width="{width}" height="70" fill="#000"/>'
        )
        svg.append(
            f'<text class="domain" x="{x+width/2}" y="{y+47}" '
            f'text-anchor="middle">{svg_escape(domain)}</text>'
        )
        midpoint = math.ceil(len(names) / 2)
        columns = [names[:midpoint], names[midpoint:]]
        col_width = (width - 60) // 2
        for col_index, table_names in enumerate(columns):
            text_x = x + 30 + col_index * col_width
            for row_index, table_name in enumerate(table_names):
                svg.append(
                    f'<text class="body" x="{text_x}" y="{y+120+row_index*46}">'
                    f'{svg_escape(table_name)}</text>'
                )
    svg.append(
        f'<text class="small" x="{PAGE_W/2}" y="{PAGE_H-95}" text-anchor="middle">'
        'Domain connectors show cross-domain foreign-key counts; detailed pages show '
        'attribute-level cardinalities.</text>'
    )
    svg.append("</svg>")
    svg_path.write_text("\n".join(svg) + "\n", encoding="utf-8")
    return png_path, svg_path


def draw_cardinality_marker(draw, point, text, toward_point):
    x, y = point
    tx, ty = toward_point
    dx = tx - x
    dy = ty - y
    if abs(dx) >= abs(dy):
        offset_x = 45 if dx > 0 else -45
        label_box(draw, (x + offset_x, y - 18), text)
    else:
        offset_y = 35 if dy > 0 else -35
        label_box(draw, (x + 45, y + offset_y), text)


def generate_domain_figure(
    schema: dict,
    tables: dict[str, dict],
    fk_columns: dict[str, set[str]],
    domain: str,
    names: list[str],
) -> tuple[Path, Path]:
    slug = DOMAIN_SLUGS[domain]
    png_path = EXPORT_DIR / f"MULTIVENT_LIVE_DATABASE_ERD_{slug}.png"
    svg_path = EXPORT_DIR / f"MULTIVENT_LIVE_DATABASE_ERD_{slug}.svg"
    positions = manuscript_layout(names)
    internal_fks = [
        fk for fk in schema["foreign_keys"]
        if fk["source_table"] in positions and fk["target_table"] in positions
    ]
    cross_fks = [
        fk for fk in schema["foreign_keys"]
        if (fk["source_table"] in positions) != (fk["target_table"] in positions)
    ]
    image = Image.new("RGB", (PAGE_W, PAGE_H), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (PAGE_W // 2, 42), f"MULTIVENT Live ERD - {domain}",
        anchor="ma", fill="black", font=FONT_TITLE,
    )
    draw.text(
        (PAGE_W // 2, 112),
        f"{len(names)} tables | {len(internal_fks)} internal foreign keys | "
        f"{len(cross_fks)} cross-domain foreign keys",
        anchor="ma", fill="black", font=FONT_SUBTITLE,
    )
    draw.text(
        (PAGE_W // 2, 172),
        "Cardinality: 1 = exactly one; 0..1 = optional one; 0..* = zero or many",
        anchor="ma", fill="black", font=FONT_SMALL,
    )
    edge_routes = []
    for edge_index, fk in enumerate(internal_fks, start=1):
        points = raster_route(
            positions[fk["source_table"]],
            positions[fk["target_table"]],
            edge_index,
        )
        edge_routes.append((fk, points))
        draw.line(points, fill="black", width=3)
    for name in names:
        table = tables[name]
        x, y, width, height = positions[name]
        draw.rectangle((x, y, x + width, y + height), fill="white", outline="black", width=4)
        draw.rectangle((x, y, x + width, y + 52), fill="black", outline="black")
        status = "RLS" if table["rls_enabled"] else "RLS OFF"
        header = f"{name}  [{status}]"
        draw.text(
            (x + 16, y + 26),
            trim_text(draw, header, FONT_HEADER, width - 32),
            anchor="lm", fill="white", font=FONT_HEADER,
        )
        lines = manuscript_column_lines(table, fk_columns[name])
        available_h = height - 66
        line_h = max(19, min(26, available_h // max(len(lines), 1)))
        for line_index, line in enumerate(lines):
            font = FONT_BODY_BOLD if line.startswith(("PK", "FK", "PK/FK")) else FONT_BODY
            draw.text(
                (x + 16, y + 65 + line_index * line_h),
                trim_text(draw, line, font, width - 32),
                fill="black", font=font,
            )
    # Draw the multiplicity labels last so table fills cannot cover them.
    for fk, points in edge_routes:
        child_mult, parent_mult = multiplicities(fk)
        draw_cardinality_marker(draw, points[0], child_mult, points[1])
        draw_cardinality_marker(draw, points[-1], parent_mult, points[-2])
    draw.text(
        (PAGE_W // 2, PAGE_H - 125),
        "PK = primary key | FK = foreign key | UQ = unique | ? = nullable | "
        "All lines use 90-degree routing",
        anchor="mm", fill="black", font=FONT_SMALL,
    )
    draw.text(
        (PAGE_W // 2, PAGE_H - 82),
        f"Source: live Supabase public schema ({PROJECT_REF}), captured {GENERATED_DATE}. "
        "See the editable Draw.io file for all 569 columns and cross-domain links.",
        anchor="mm", fill="black", font=FONT_SMALL,
    )
    image.save(png_path, dpi=(300, 300))

    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{PAGE_W}" height="{PAGE_H}" '
        f'viewBox="0 0 {PAGE_W} {PAGE_H}">',
        '<rect width="100%" height="100%" fill="white"/>',
        '<style>text{font-family:Arial,sans-serif;fill:#000}'
        '.title{font-size:48px;font-weight:700}.subtitle{font-size:27px}'
        '.header{font-size:24px;font-weight:700;fill:#fff}'
        '.body{font-size:19px}.key{font-size:19px;font-weight:700}'
        '.small{font-size:17px}</style>',
        f'<text class="title" x="{PAGE_W/2}" y="82" text-anchor="middle">'
        f'MULTIVENT Live ERD - {svg_escape(domain)}</text>',
        f'<text class="subtitle" x="{PAGE_W/2}" y="135" text-anchor="middle">'
        f'{len(names)} tables | {len(internal_fks)} internal foreign keys | '
        f'{len(cross_fks)} cross-domain foreign keys</text>',
        f'<text class="small" x="{PAGE_W/2}" y="188" text-anchor="middle">'
        'Cardinality: 1 = exactly one; 0..1 = optional one; 0..* = zero or many</text>',
    ]
    for fk, points in edge_routes:
        point_text = " ".join(f"{x},{y}" for x, y in points)
        svg.append(
            f'<polyline points="{point_text}" fill="none" stroke="#000" stroke-width="3"/>'
        )
    for name in names:
        table = tables[name]
        x, y, width, height = positions[name]
        status = "RLS" if table["rls_enabled"] else "RLS OFF"
        svg.append(
            f'<rect x="{x}" y="{y}" width="{width}" height="{height}" '
            'fill="#fff" stroke="#000" stroke-width="4"/>'
        )
        svg.append(
            f'<rect x="{x}" y="{y}" width="{width}" height="52" fill="#000"/>'
        )
        svg.append(
            f'<text class="header" x="{x+16}" y="{y+35}">'
            f'{svg_escape(name)} [{status}]</text>'
        )
        lines = manuscript_column_lines(table, fk_columns[name])
        available_h = height - 66
        line_h = max(19, min(26, available_h // max(len(lines), 1)))
        for line_index, line in enumerate(lines):
            css = "key" if line.startswith(("PK", "FK", "PK/FK")) else "body"
            svg.append(
                f'<text class="{css}" x="{x+16}" y="{y+82+line_index*line_h}">'
                f'{svg_escape(line)}</text>'
            )
    # Cardinality labels sit above the table shapes in SVG stacking order.
    for fk, points in edge_routes:
        child_mult, parent_mult = multiplicities(fk)
        for point, toward, value in (
            (points[0], points[1], child_mult),
            (points[-1], points[-2], parent_mult),
        ):
            px, py = point
            tx, ty = toward
            if abs(tx - px) >= abs(ty - py):
                lx = px + (45 if tx > px else -45)
                ly = py - 18
            else:
                lx = px + 45
                ly = py + (35 if ty > py else -35)
            width = 52 if value != "0..*" else 64
            svg.append(
                f'<rect x="{lx-width/2}" y="{ly-14}" width="{width}" height="28" '
                'fill="#fff" stroke="#000"/>'
            )
            svg.append(
                f'<text class="small" x="{lx}" y="{ly+6}" text-anchor="middle">'
                f'{svg_escape(value)}</text>'
            )
    svg.append(
        f'<text class="small" x="{PAGE_W/2}" y="{PAGE_H-120}" text-anchor="middle">'
        'PK = primary key | FK = foreign key | UQ = unique | ? = nullable | '
        'All lines use 90-degree routing</text>'
    )
    svg.append(
        f'<text class="small" x="{PAGE_W/2}" y="{PAGE_H-78}" text-anchor="middle">'
        f'Source: live Supabase public schema ({PROJECT_REF}), captured {GENERATED_DATE}. '
        'See the editable Draw.io file for all attributes and cross-domain links.</text>'
    )
    svg.append("</svg>")
    svg_path.write_text("\n".join(svg) + "\n", encoding="utf-8")
    return png_path, svg_path


def build_pdf(png_paths: list[Path]) -> None:
    images = [Image.open(path).convert("RGB") for path in png_paths]
    first, rest = images[0], images[1:]
    first.save(
        PDF_PATH,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=rest,
    )
    for image in images:
        image.close()


def build_docx(png_paths: list[Path], captions: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for index, (image_path, caption) in enumerate(zip(png_paths, captions)):
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(caption)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(12)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(10.8))
        source = document.add_paragraph(
            f"Source: MULTIVENT live Supabase public schema, captured {GENERATED_DATE}."
        )
        source.alignment = WD_ALIGN_PARAGRAPH.CENTER
        source.runs[0].italic = True
        source.runs[0].font.size = Pt(8)
        if index < len(png_paths) - 1:
            document.add_page_break()
    document.save(DOCX_PATH)


def write_notes(schema: dict, generated_files: list[Path]) -> None:
    tables = table_map(schema)
    fk_columns = fk_column_map(schema)
    lines = [
        "# MULTIVENT Live Database ERD Export",
        "",
        f"- Supabase project reference: `{PROJECT_REF}`",
        f"- Schema: `public`",
        f"- Captured: {GENERATED_DATE}",
        f"- Tables: {len(schema['tables'])}",
        f"- Columns: {sum(len(t['columns']) for t in schema['tables'])}",
        f"- Foreign keys: {len(schema['foreign_keys'])}",
        "",
        "## Files",
        "",
    ]
    for path in generated_files:
        lines.append(f"- `{path.name}`")
    lines += [
        "",
        "## Diagram conventions",
        "",
        "- PK: primary key",
        "- FK: foreign key",
        "- UQ: unique column",
        "- ?: nullable column",
        "- 1: exactly one parent",
        "- 0..1: optional parent or optional unique child",
        "- 0..*: zero or many child rows",
        "- All connectors use orthogonal, 90-degree routing.",
        "",
        "## Domain inventory",
        "",
    ]
    for domain, names in DOMAIN_TABLES.items():
        internal = [
            fk for fk in schema["foreign_keys"]
            if fk["source_table"] in names and fk["target_table"] in names
        ]
        lines.append(f"### {domain}")
        lines.append("")
        lines.append(
            f"{len(names)} tables and {len(internal)} internal foreign-key relationships."
        )
        lines.append("")
        for name in names:
            lines.append(
                f"- `{name}`: {len(tables[name]['columns'])} columns; "
                f"{len(fk_columns[name])} FK columns"
            )
        lines.append("")
    lines += [
        "## Important live-schema note",
        "",
        "`service_categories` had RLS disabled when this export was captured. "
        "The ERD records the deployed state and does not change database security.",
        "",
    ]
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    schema = load_schema()
    tables = table_map(schema)
    fk_columns = fk_column_map(schema)

    create_drawio(schema)
    overview_png, overview_svg = generate_overview(schema)
    png_paths = [overview_png]
    svg_paths = [overview_svg]
    captions = ["Figure 1. Condensed MULTIVENT live database ERD overview"]
    for index, (domain, names) in enumerate(DOMAIN_TABLES.items(), start=2):
        png_path, svg_path = generate_domain_figure(
            schema, tables, fk_columns, domain, names
        )
        png_paths.append(png_path)
        svg_paths.append(svg_path)
        captions.append(f"Figure {index}. MULTIVENT ERD - {domain}")

    build_pdf(png_paths)
    build_docx(png_paths, captions)
    generated = [
        DRAWIO_PATH, PDF_PATH, DOCX_PATH, *png_paths, *svg_paths, NOTES_PATH
    ]
    write_notes(schema, generated)
    print(
        json.dumps(
            {
                "tables": len(schema["tables"]),
                "columns": sum(len(t["columns"]) for t in schema["tables"]),
                "foreign_keys": len(schema["foreign_keys"]),
                "files": [str(path) for path in generated],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()





