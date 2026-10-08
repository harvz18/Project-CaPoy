from __future__ import annotations

import math
from collections import Counter
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt

from generate_live_database_erd import (
    DOMAIN_SLUGS,
    DOMAIN_TABLES,
    EXPORT_DIR,
    compact_type,
    domain_for_table,
    fk_column_map,
    load_schema,
    table_map,
)


OUTPUT_DIR = EXPORT_DIR / "white_large"
WIDTH = 3508
HEIGHT = 2480
TABLES_PER_PAGE = 6
PDF_PATH = OUTPUT_DIR / "MULTIVENT_ERD_WHITE_LARGE_MANUSCRIPT.pdf"
DOCX_PATH = OUTPUT_DIR / "MULTIVENT_ERD_WHITE_LARGE_MANUSCRIPT.docx"

WHITE = (255, 255, 255)
GRID = (226, 226, 226)
BLACK = (0, 0, 0)
DARK = (25, 25, 25)
MID = (92, 92, 92)
BORDER = (45, 45, 45)
LINE = (78, 78, 78)
ROW_A = (255, 255, 255)
ROW_B = (245, 245, 245)
HEADER = (28, 28, 28)


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


FONT_TITLE = font(56, True)
FONT_SUBTITLE = font(30)
FONT_DOMAIN = font(32, True)
FONT_HEADER = font(34, True)
FONT_COLUMN = font(31)
FONT_COLUMN_BOLD = font(31, True)
FONT_TYPE = font(26)
FONT_SMALL = font(24)
FONT_TINY = font(23)


def trim(draw: ImageDraw.ImageDraw, value: str, text_font, max_width: int) -> str:
    if draw.textlength(value, font=text_font) <= max_width:
        return value
    text = value
    while text and draw.textlength(text + "...", font=text_font) > max_width:
        text = text[:-1]
    return text.rstrip() + "..."


def draw_grid(draw: ImageDraw.ImageDraw) -> None:
    spacing = 42
    for y in range(20, HEIGHT, spacing):
        for x in range(20, WIDTH, spacing):
            draw.ellipse((x - 1, y - 1, x + 1, y + 1), fill=GRID)


def is_fk(table_name: str, column_name: str, fk_columns) -> bool:
    return column_name in fk_columns.get(table_name, set())


def key_kind(table_name: str, column: dict, fk_columns) -> str:
    foreign = is_fk(table_name, column["name"], fk_columns)
    if column["primary_key"] and foreign:
        return "PK/FK"
    if column["primary_key"]:
        return "PK"
    if foreign:
        return "FK"
    if column["unique"]:
        return "UQ"
    return ""


def draw_table_icon(draw: ImageDraw.ImageDraw, x: int, y: int, light: bool = False):
    color = WHITE if light else BLACK
    draw.rectangle((x, y, x + 25, y + 25), outline=color, width=2)
    draw.line((x, y + 8, x + 25, y + 8), fill=color, width=2)
    draw.line((x + 8, y, x + 8, y + 25), fill=color, width=2)


def draw_column_icon(draw: ImageDraw.ImageDraw, x: int, y: int, kind: str):
    if kind in {"PK", "PK/FK"}:
        draw.ellipse((x - 7, y - 7, x + 7, y + 7), outline=BLACK, width=3)
        draw.line((x + 7, y, x + 22, y), fill=BLACK, width=3)
        draw.line((x + 15, y, x + 15, y + 7), fill=BLACK, width=3)
        draw.line((x + 22, y, x + 22, y + 6), fill=BLACK, width=3)
    elif kind in {"FK", "UQ"}:
        draw.polygon(
            [(x, y - 8), (x + 8, y), (x, y + 8), (x - 8, y)],
            outline=BLACK,
        )
    else:
        draw.ellipse((x - 6, y - 6, x + 6, y + 6), outline=MID, width=2)


def chunked(values, size):
    return [values[index:index + size] for index in range(0, len(values), size)]


def page_layout(names):
    columns = 3
    rows = math.ceil(len(names) / columns)
    margin_x = 100
    gap_x = 55
    top = 230
    bottom = 115
    gap_y = 80
    card_w = (WIDTH - 2 * margin_x - gap_x * (columns - 1)) // columns
    available_h = HEIGHT - top - bottom - (rows - 1) * gap_y
    card_h = min(940, available_h // rows)
    block_h = rows * card_h + (rows - 1) * gap_y
    start_y = top + max(0, (HEIGHT - top - bottom - block_h) // 2)
    positions = {}
    for index, name in enumerate(names):
        row, column = divmod(index, columns)
        positions[name] = (
            margin_x + column * (card_w + gap_x),
            start_y + row * (card_h + gap_y),
            card_w,
            card_h,
        )
    return positions


def displayed_columns(table_name, table, fk_columns, card_h):
    header_h = 70
    row_h = 47
    max_rows = max(5, (card_h - header_h - 12) // row_h)
    columns = table["columns"]
    if len(columns) <= max_rows:
        return columns, row_h
    required = {
        column["name"]
        for column in columns
        if column["primary_key"] or is_fk(table_name, column["name"], fk_columns)
    }
    selected_names = set(required)
    target_count = max(max_rows - 1, len(required))
    for column in columns:
        if len(selected_names) >= target_count:
            break
        selected_names.add(column["name"])
    selected = [column for column in columns if column["name"] in selected_names]
    omitted = len(columns) - len(selected)
    if omitted:
        selected.append(
            {
                "name": f"... +{omitted} additional columns",
                "type": "",
                "nullable": True,
                "primary_key": False,
                "unique": False,
                "_summary": True,
            }
        )
    return selected[:max_rows], row_h


def prepare_cards(names, tables, fk_columns, positions):
    cards = {}
    for name in names:
        box = positions[name]
        columns, row_h = displayed_columns(name, tables[name], fk_columns, box[3])
        row_map = {
            column["name"]: box[1] + 70 + index * row_h + row_h // 2
            for index, column in enumerate(columns)
            if not column.get("_summary")
        }
        cards[name] = {
            "box": box,
            "columns": columns,
            "row_h": row_h,
            "row_map": row_map,
        }
    return cards


def route(source, target, source_column, target_column, edge_index):
    sx, sy, sw, sh = source["box"]
    tx, ty, tw, th = target["box"]
    source_y = source["row_map"].get(source_column, sy + 96)
    target_y = target["row_map"].get(target_column, ty + 96)
    if sx + sw < tx:
        start = (sx + sw, source_y)
        end = (tx, target_y)
        middle_x = (start[0] + end[0]) // 2 + ((edge_index % 7) - 3) * 6
        return [start, (middle_x, source_y), (middle_x, target_y), end]
    if tx + tw < sx:
        start = (sx, source_y)
        end = (tx + tw, target_y)
        middle_x = (start[0] + end[0]) // 2 + ((edge_index % 7) - 3) * 6
        return [start, (middle_x, source_y), (middle_x, target_y), end]
    right_side = edge_index % 2 == 0
    if right_side:
        start = (sx + sw, source_y)
        end = (tx + tw, target_y)
        outer_x = max(sx + sw, tx + tw) + 34 + (edge_index % 7) * 10
    else:
        start = (sx, source_y)
        end = (tx, target_y)
        outer_x = min(sx, tx) - 34 - (edge_index % 7) * 10
    return [start, (outer_x, source_y), (outer_x, target_y), end]


def draw_table(draw, table_name, table, card, fk_columns):
    x, y, width, height = card["box"]
    draw.rectangle((x, y, x + width, y + height), fill=WHITE, outline=BORDER, width=3)
    draw.rectangle((x, y, x + width, y + 70), fill=HEADER, outline=HEADER)
    draw_table_icon(draw, x + 18, y + 22, light=True)
    draw.text(
        (x + 58, y + 35),
        trim(draw, table_name, FONT_HEADER, width - 105),
        anchor="lm",
        fill=WHITE,
        font=FONT_HEADER,
    )
    row_h = card["row_h"]
    for index, column in enumerate(card["columns"]):
        row_y = y + 70 + index * row_h
        draw.rectangle(
            (x + 1, row_y, x + width - 1, row_y + row_h),
            fill=ROW_B if index % 2 else ROW_A,
        )
        draw.line((x, row_y + row_h, x + width, row_y + row_h), fill=GRID, width=2)
        if column.get("_summary"):
            draw.text(
                (x + 22, row_y + row_h // 2),
                column["name"],
                anchor="lm",
                fill=MID,
                font=FONT_SMALL,
            )
            continue
        kind = key_kind(table_name, column, fk_columns)
        draw_column_icon(draw, x + 27, row_y + row_h // 2, kind)
        column_font = FONT_COLUMN_BOLD if kind else FONT_COLUMN
        nullable = " ?" if column["nullable"] else ""
        draw.text(
            (x + 64, row_y + row_h // 2),
            trim(draw, column["name"] + nullable, column_font, int(width * 0.60) - 70),
            anchor="lm",
            fill=BLACK,
            font=column_font,
        )
        draw.text(
            (x + width - 20, row_y + row_h // 2),
            trim(draw, compact_type(column["type"]), FONT_TYPE, int(width * 0.32)),
            anchor="rm",
            fill=MID,
            font=FONT_TYPE,
        )


def page_title(domain, part_number, total_parts):
    if total_parts == 1:
        return f"MULTIVENT ERD - {domain}"
    return f"MULTIVENT ERD - {domain} ({part_number} of {total_parts})"


def generate_detail_page(schema, tables, fk_columns, domain, names, part_number, total_parts):
    positions = page_layout(names)
    cards = prepare_cards(names, tables, fk_columns, positions)
    relationships = [
        fk for fk in schema["foreign_keys"]
        if fk["source_table"] in cards and fk["target_table"] in cards
    ]
    image = Image.new("RGB", (WIDTH, HEIGHT), WHITE)
    draw = ImageDraw.Draw(image)
    draw_grid(draw)
    draw.text(
        (WIDTH // 2, 40),
        page_title(domain, part_number, total_parts),
        anchor="ma",
        fill=BLACK,
        font=FONT_TITLE,
    )
    draw.text(
        (WIDTH // 2, 112),
        f"{len(names)} tables  |  {len(relationships)} relationships",
        anchor="ma",
        fill=MID,
        font=FONT_SUBTITLE,
    )

    routes = []
    for edge_index, relationship in enumerate(relationships, start=1):
        points = route(
            cards[relationship["source_table"]],
            cards[relationship["target_table"]],
            relationship["source_columns"][0],
            relationship["target_columns"][0],
            edge_index,
        )
        routes.append((relationship, points))
        draw.line(points, fill=LINE, width=3)

    for name in names:
        draw_table(draw, name, tables[name], cards[name], fk_columns)

    for relationship, points in routes:
        draw.ellipse(
            (points[0][0] - 4, points[0][1] - 4, points[0][0] + 4, points[0][1] + 4),
            fill=BLACK,
        )
        draw.ellipse(
            (points[-1][0] - 4, points[-1][1] - 4, points[-1][0] + 4, points[-1][1] + 4),
            fill=BLACK,
        )
    draw.text(
        (72, HEIGHT - 54),
        "PK key  |  FK/UQ diamond  |  ? nullable",
        anchor="lm",
        fill=MID,
        font=FONT_SMALL,
    )
    suffix = f"_PART_{part_number}" if total_parts > 1 else ""
    path = OUTPUT_DIR / (
        f"MULTIVENT_ERD_WHITE_LARGE_{DOMAIN_SLUGS[domain]}{suffix}.png"
    )
    image.save(path, dpi=(300, 300))
    return path


def panel_route(source_box, target_box, index):
    sx, sy, sw, sh = source_box
    tx, ty, tw, th = target_box
    source_center = (sx + sw // 2, sy + sh // 2)
    target_center = (tx + tw // 2, ty + th // 2)
    if abs(target_center[0] - source_center[0]) > abs(target_center[1] - source_center[1]):
        if target_center[0] > source_center[0]:
            start, end = (sx + sw, source_center[1]), (tx, target_center[1])
        else:
            start, end = (sx, source_center[1]), (tx + tw, target_center[1])
        middle = (start[0] + end[0]) // 2 + ((index % 5) - 2) * 7
        return [start, (middle, start[1]), (middle, end[1]), end]
    if target_center[1] > source_center[1]:
        start, end = (source_center[0], sy + sh), (target_center[0], ty)
    else:
        start, end = (source_center[0], sy), (target_center[0], ty + th)
    middle = (start[1] + end[1]) // 2 + ((index % 5) - 2) * 7
    return [start, (start[0], middle), (end[0], middle), end]


def generate_overview(schema):
    image = Image.new("RGB", (WIDTH, HEIGHT), WHITE)
    draw = ImageDraw.Draw(image)
    draw_grid(draw)
    draw.text(
        (WIDTH // 2, 40),
        "MULTIVENT ERD - Overview",
        anchor="ma",
        fill=BLACK,
        font=FONT_TITLE,
    )
    draw.text(
        (WIDTH // 2, 112),
        f"{len(schema['tables'])} tables  |  "
        f"{sum(len(table['columns']) for table in schema['tables'])} columns  |  "
        f"{len(schema['foreign_keys'])} foreign keys",
        anchor="ma",
        fill=MID,
        font=FONT_SUBTITLE,
    )
    positions = {
        "Identity and Access": (55, 235, 1000, 690),
        "Marketplace and Availability": (1254, 235, 1000, 780),
        "Booking and Finance": (2453, 235, 1000, 640),
        "Event Planning and Coordination": (310, 1390, 1300, 850),
        "Communication, Feedback, and Operations": (1898, 1390, 1300, 850),
    }
    counts = Counter()
    for relationship in schema["foreign_keys"]:
        source = domain_for_table(relationship["source_table"])
        target = domain_for_table(relationship["target_table"])
        if source != target:
            counts[(source, target)] += 1
    for index, ((source, target), count) in enumerate(sorted(counts.items()), start=1):
        points = panel_route(positions[source], positions[target], index)
        draw.line(points, fill=LINE, width=3)
        a, b = points[1], points[2]
        lx, ly = (a[0] + b[0]) // 2, (a[1] + b[1]) // 2
        draw.rectangle((lx - 50, ly - 21, lx + 50, ly + 21), fill=WHITE, outline=BORDER, width=2)
        draw.text((lx, ly), f"{count} FK", anchor="mm", fill=BLACK, font=FONT_TINY)

    for domain, names in DOMAIN_TABLES.items():
        x, y, width, height = positions[domain]
        draw.rectangle((x, y, x + width, y + height), fill=WHITE, outline=BORDER, width=3)
        draw.rectangle((x, y, x + width, y + 64), fill=HEADER)
        draw_table_icon(draw, x + 18, y + 20, light=True)
        draw.text((x + 58, y + 32), domain, anchor="lm", fill=WHITE, font=FONT_DOMAIN)
        columns = 2
        gap = 16
        mini_w = (width - 44 - gap) // 2
        mini_h = 49
        for table_index, name in enumerate(names):
            row, column = divmod(table_index, columns)
            bx = x + 18 + column * (mini_w + gap)
            by = y + 88 + row * (mini_h + 12)
            draw.rectangle(
                (bx, by, bx + mini_w, by + mini_h),
                fill=ROW_B,
                outline=GRID,
                width=2,
            )
            draw_table_icon(draw, bx + 11, by + 12)
            draw.text(
                (bx + 48, by + mini_h // 2),
                trim(draw, name, FONT_TINY, mini_w - 62),
                anchor="lm",
                fill=BLACK,
                font=FONT_TINY,
            )
    path = OUTPUT_DIR / "MULTIVENT_ERD_WHITE_LARGE_OVERVIEW.png"
    image.save(path, dpi=(300, 300))
    return path


def build_pdf(paths):
    images = [Image.open(path).convert("RGB") for path in paths]
    images[0].save(
        PDF_PATH,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=images[1:],
    )
    for image in images:
        image.close()


def build_docx(paths, captions):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(8)
    section.bottom_margin = Mm(8)
    section.left_margin = Mm(8)
    section.right_margin = Mm(8)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)
    for index, (path, caption) in enumerate(zip(paths, captions)):
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(caption)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(12)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(10.9))
        if index < len(paths) - 1:
            document.add_page_break()
    document.save(DOCX_PATH)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for previous_image in OUTPUT_DIR.glob("MULTIVENT_ERD_WHITE_LARGE_*.png"):
        previous_image.unlink()
    schema = load_schema()
    tables = table_map(schema)
    fk_columns = fk_column_map(schema)
    paths = [generate_overview(schema)]
    captions = ["MULTIVENT ERD - Overview"]
    for domain, names in DOMAIN_TABLES.items():
        parts = chunked(names, TABLES_PER_PAGE)
        for part_number, part in enumerate(parts, start=1):
            paths.append(
                generate_detail_page(
                    schema,
                    tables,
                    fk_columns,
                    domain,
                    part,
                    part_number,
                    len(parts),
                )
            )
            captions.append(page_title(domain, part_number, len(parts)))
    build_pdf(paths)
    build_docx(paths, captions)
    print(f"Generated {len(paths)} large-font white ERD images.")
    for path in paths:
        print(path)
    print(PDF_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()

