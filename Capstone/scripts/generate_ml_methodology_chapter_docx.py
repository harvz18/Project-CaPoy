r"""Generate the manuscript-ready MULTIVENT Section 3.2 ML methodology DOCX.

The content is intentionally evidence-based. Counts, parameters, and metrics
come from the generated artifacts in D:\Sentiment Training\multivent-sentiment.
"""

from __future__ import annotations

from pathlib import Path
from copy import deepcopy

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "3.2_MACHINE_LEARNING_METHODOLOGY_SENTIMENT_MODEL_REVISED.docx"
ASSET_DIR = ROOT / ".temp" / "ml_methodology_chapter"
PIPELINE_IMAGE = ASSET_DIR / "figure_3_2_1_training_pipeline.png"
RUNTIME_IMAGE = ASSET_DIR / "figure_3_2_2_runtime_pipeline.png"
MATHML_TO_OMML_XSL = Path(
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
)


# Presentation MathML is transformed into native Office Math Markup Language
# when the manuscript is generated. The resulting equations remain editable in
# Microsoft Word and use built-up fractions, limits, accents, and subscripts.
EQUATION_MATHML = {
    1: """
      <mrow><mover accent="true"><mi>v</mi><mo>^</mo></mover><mo>=</mo>
      <mfrac><mi>v</mi><mrow><mi mathvariant="normal">max</mi><mo>(</mo>
      <msub><mrow><mo>‖</mo><mi>v</mi><mo>‖</mo></mrow><mn>2</mn></msub>
      <mo>,</mo><mi>ε</mi><mo>)</mo></mrow></mfrac></mrow>""",
    2: """
      <mrow><msub><mi>K</mi><mi>i</mi></msub><mo>=</mo>
      <mi mathvariant="normal">min</mi><mo>(</mo><mn>1</mn><mo>,</mo>
      <mfrac><msub><mi>h</mi><mi>i</mi></msub><mn>2</mn></mfrac><mo>)</mo></mrow>""",
    3: """
      <mrow><msub><mi>R</mi><mi>i</mi></msub><mo>=</mo>
      <msub><mi mathvariant="normal">clip</mi><mrow><mo>[</mo><mn>0</mn><mo>,</mo><mn>1</mn><mo>]</mo></mrow></msub>
      <mo>(</mo><mn>0.55</mn><msub><mi>S</mi><mi>i</mi></msub><mo>+</mo>
      <mn>0.30</mn><msub><mi>K</mi><mi>i</mi></msub><mo>+</mo>
      <mn>0.15</mn><msub><mi>C</mi><mi>i</mi></msub><mo>−</mo>
      <msub><mi>P</mi><mi>i</mi></msub><mo>)</mo></mrow>""",
    4: """
      <mrow><msub><mi>θ</mi><mi>d</mi></msub><mo>∼</mo>
      <mi mathvariant="normal">Dirichlet</mi><mo>(</mo><mi>α</mi><mo>)</mo><mo>,</mo>
      <mspace width="1em"/><msub><mi>φ</mi><mi>k</mi></msub><mo>∼</mo>
      <mi mathvariant="normal">Dirichlet</mi><mo>(</mo><mi>β</mi><mo>)</mo></mrow>""",
    5: """
      <mrow><msub><mi>z</mi><mrow><mi>d</mi><mi>n</mi></mrow></msub><mo>∼</mo>
      <mi mathvariant="normal">Categorical</mi><mo>(</mo><msub><mi>θ</mi><mi>d</mi></msub><mo>)</mo><mo>,</mo>
      <mspace width="1em"/><msub><mi>w</mi><mrow><mi>d</mi><mi>n</mi></mrow></msub><mo>∼</mo>
      <mi mathvariant="normal">Categorical</mi><mo>(</mo>
      <msub><mi>φ</mi><msub><mi>z</mi><mrow><mi>d</mi><mi>n</mi></mrow></msub></msub><mo>)</mo></mrow>""",
    6: """
      <mrow><msup><mi>k</mi><mo>*</mo></msup><mo>=</mo>
      <munder><mi mathvariant="normal">arg max</mi><mi>k</mi></munder>
      <mi>P</mi><mo>(</mo><mi mathvariant="normal">topic</mi><mo>=</mo><mi>k</mi>
      <mo>|</mo><mi mathvariant="normal">document</mi><mspace width="0.3em"/><mi>d</mi><mo>)</mo></mrow>""",
    7: """
      <mrow><mi mathvariant="normal">sim</mi><mo>(</mo><mi>q</mi><mo>,</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo><mo>=</mo>
      <msup><mover accent="true"><mi>q</mi><mo>^</mo></mover><mi>T</mi></msup>
      <msub><mover accent="true"><mi>x</mi><mo>^</mo></mover><mi>i</mi></msub><mo>=</mo>
      <mi mathvariant="normal">cos</mi><mo>(</mo><mi>q</mi><mo>,</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo></mrow>""",
    8: """
      <mrow><msub><mi>W</mi><mi>i</mi></msub><mo>=</mo>
      <mi mathvariant="normal">max</mi><mo>(</mo><mn>0</mn><mo>,</mo>
      <mi mathvariant="normal">sim</mi><mo>(</mo><mi>q</mi><mo>,</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo><mo>)</mo>
      <mo>×</mo><msub><mi>R</mi><mi>i</mi></msub><mo>×</mo>
      <msub><mi>ω</mi><mi mathvariant="normal">source</mi></msub></mrow>""",
    9: """
      <mrow><msub><mi>E</mi><mi>c</mi></msub><mo>=</mo>
      <munder><mo>∑</mo><mrow><mi>i</mi><mo>:</mo><msub><mi>y</mi><mi>i</mi></msub><mo>=</mo><mi>c</mi></mrow></munder>
      <msub><mi>W</mi><mi>i</mi></msub><mo>,</mo><mspace width="1em"/>
      <msub><mi>p</mi><mi>c</mi></msub><mo>=</mo>
      <mfrac><msub><mi>E</mi><mi>c</mi></msub><mrow>
      <msub><mi>E</mi><mi mathvariant="normal">positive</mi></msub><mo>+</mo>
      <msub><mi>E</mi><mi mathvariant="normal">negative</mi></msub></mrow></mfrac></mrow>""",
    10: """
      <mrow><msub><mi>s</mi><mi mathvariant="normal">sentiment</mi></msub><mo>=</mo>
      <msub><mi>p</mi><mi mathvariant="normal">positive</mi></msub><mo>−</mo>
      <msub><mi>p</mi><mi mathvariant="normal">negative</mi></msub><mo>,</mo><mspace width="1em"/>
      <mo>−</mo><mn>1</mn><mo>≤</mo><msub><mi>s</mi><mi mathvariant="normal">sentiment</mi></msub><mo>≤</mo><mn>1</mn></mrow>""",
    11: """
      <mrow><msub><mi>u</mi><mi>g</mi></msub><mo>=</mo>
      <mfrac><mrow><mi mathvariant="normal">int</mi><mo>(</mo>
      <msub><mi mathvariant="normal">SHA256</mi><mrow><mo>[</mo><mn>0</mn><mo>:</mo><mn>16</mn><mo>]</mo></mrow></msub>
      <mo>(</mo><mtext>“42:”</mtext><mo>∥</mo><msub><mi mathvariant="normal">group_id</mi><mi>g</mi></msub><mo>)</mo>
      <mo>,</mo><mn>16</mn><mo>)</mo></mrow><msup><mn>2</mn><mn>64</mn></msup></mfrac></mrow>""",
    12: """
      <mrow><mi mathvariant="normal">Accuracy</mi><mo>=</mo>
      <mfrac><mrow><mi>TP</mi><mo>+</mo><mi>TN</mi></mrow>
      <mrow><mi>TP</mi><mo>+</mo><mi>TN</mi><mo>+</mo><mi>FP</mi><mo>+</mo><mi>FN</mi></mrow></mfrac></mrow>""",
    13: """
      <mrow><mi mathvariant="normal">Precision</mi><mo>=</mo>
      <mfrac><mi>TP</mi><mrow><mi>TP</mi><mo>+</mo><mi>FP</mi></mrow></mfrac>
      <mo>,</mo><mspace width="1.5em"/><mi mathvariant="normal">Recall</mi><mo>=</mo>
      <mfrac><mi>TP</mi><mrow><mi>TP</mi><mo>+</mo><mi>FN</mi></mrow></mfrac></mrow>""",
    14: """
      <mrow><msub><mi>F</mi><mn>1</mn></msub><mo>=</mo><mn>2</mn>
      <mfrac><mrow><mi mathvariant="normal">Precision</mi><mo>×</mo><mi mathvariant="normal">Recall</mi></mrow>
      <mrow><mi mathvariant="normal">Precision</mi><mo>+</mo><mi mathvariant="normal">Recall</mi></mrow></mfrac></mrow>""",
    15: """
      <mrow><msubsup><mi>F</mi><mn>1</mn><mi mathvariant="normal">macro</mi></msubsup><mo>=</mo>
      <mfrac><mrow><msubsup><mi>F</mi><mn>1</mn><mi mathvariant="normal">positive</mi></msubsup><mo>+</mo>
      <msubsup><mi>F</mi><mn>1</mn><mi mathvariant="normal">negative</mi></msubsup></mrow><mn>2</mn></mfrac></mrow>""",
}


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, *, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def add_caption(document: Document, text: str, *, above=True) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = above
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4 if above else 8)
    run = paragraph.add_run(text)
    run.bold = True


def add_table(document: Document, caption: str, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None, font_size=9.0) -> None:
    add_caption(document, caption, above=True)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, size=font_size,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[index], "E7E7E7")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            set_cell_text(cells[column_index], value, size=font_size)
            if row_index % 2 == 1:
                set_cell_shading(cells[column_index], "F6F6F6")
    if widths:
        set_table_widths(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_equation(document: Document, equation: str, number: int,
                 definition: str | None = None) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.25)
    table.columns[1].width = Inches(0.55)
    formula_cell, number_cell = table.rows[0].cells
    formula_cell.text = ""
    formula_paragraph = formula_cell.paragraphs[0]
    formula_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula_paragraph.paragraph_format.space_before = Pt(4)
    formula_paragraph.paragraph_format.space_after = Pt(4)
    formula_paragraph.paragraph_format.keep_together = True
    formula_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(formula_cell, top=100, bottom=100)

    if MATHML_TO_OMML_XSL.exists() and number in EQUATION_MATHML:
        mathml = (
            '<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">'
            + EQUATION_MATHML[number]
            + "</math>"
        )
        transform = etree.XSLT(etree.parse(str(MATHML_TO_OMML_XSL)))
        omml_root = transform(etree.fromstring(mathml.encode("utf-8"))).getroot()
        formula_paragraph._p.append(deepcopy(omml_root))
    else:
        # Portable fallback for systems that do not have the Microsoft Office
        # MathML transform. The distributed DOCX is generated with native OMML.
        run = formula_paragraph.add_run(equation)
        run.font.name = "Cambria Math"
        run.font.size = Pt(11.5)
    set_cell_text(number_cell, f"({number})", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for cell in (formula_cell, number_cell):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "nil")
            borders.append(node)
        tc_pr.append(borders)
    if definition:
        paragraph = document.add_paragraph(definition)
        paragraph.style = document.styles["Equation Explanation"]


def add_body(document: Document, text: str, *, first_line=True) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.4) if first_line else None
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.25
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def draw_box(ax, x, y, w, h, title, body=""):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.01,rounding_size=0.02",
        linewidth=1.35, edgecolor="black", facecolor="white",
    )
    ax.add_patch(patch)
    label = title if not body else f"{title}\n{body}"
    ax.text(x + w / 2, y + h / 2, label, ha="center", va="center",
            fontsize=8.3, fontweight="bold" if not body else "normal", wrap=True)


def arrow(ax, start, end):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=10,
                                 linewidth=1.15, color="black"))


def create_training_pipeline(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7.4, 8.2), dpi=260)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 13)
    ax.axis("off")
    stages = [
        (2.1, 11.5, 5.8, 0.9, "Registered data sources", "HiliSenti · MAMS · hotel reviews · UCI"),
        (2.1, 10.0, 5.8, 0.9, "Schema and label validation", "Binary mapping; neutral/invalid units excluded"),
        (2.1, 8.5, 5.8, 0.9, "Conservative text preprocessing", "Unicode, whitespace, informal normalization, deduplication"),
        (2.1, 7.0, 5.8, 0.9, "Fixed data partitioning", "Training · validation · test; group and leakage guards"),
        (2.1, 5.5, 5.8, 0.9, "Domain relevance filtering", "Frozen embeddings + keywords + category prior"),
    ]
    for item in stages:
        draw_box(ax, *item)
    for y1, y2 in ((11.5, 10.9), (10.0, 9.4), (8.5, 7.9), (7.0, 6.4)):
        arrow(ax, (5, y1), (5, y2))
    draw_box(ax, 0.7, 3.6, 4.0, 1.1, "LDA training branch", "4,462 training rows; 10 topics; count features")
    draw_box(ax, 5.3, 3.6, 4.0, 1.1, "Retrieval branch", "3,050 balanced rows; 384-D embeddings; FAISS")
    ax.plot([5, 5, 2.7], [5.5, 5.05, 5.05], color="black", linewidth=1.1)
    arrow(ax, (2.7, 5.05), (2.7, 4.7))
    ax.plot([5, 5, 7.3], [5.5, 5.05, 5.05], color="black", linewidth=1.1)
    arrow(ax, (7.3, 5.05), (7.3, 4.7))
    draw_box(ax, 0.7, 1.7, 4.0, 1.0, "Saved LDA artifacts", "vectorizer.joblib · lda_model.joblib · topics.json")
    draw_box(ax, 5.3, 1.7, 4.0, 1.0, "Saved retrieval artifacts", "index.faiss · metadata.json · manifest.json")
    arrow(ax, (2.7, 3.6), (2.7, 2.7))
    arrow(ax, (7.3, 3.6), (7.3, 2.7))
    draw_box(ax, 2.1, 0.2, 5.8, 0.8, "Held-out evaluation", "Fixed test samples; accuracy, precision, recall, F1, macro F1")
    ax.plot([2.7, 2.7, 5], [1.7, 1.35, 1.35], color="black", linewidth=1.1)
    ax.plot([7.3, 7.3, 5], [1.7, 1.35, 1.35], color="black", linewidth=1.1)
    arrow(ax, (5, 1.35), (5, 1.0))
    fig.tight_layout(pad=0.4)
    fig.savefig(path, dpi=260, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def create_runtime_pipeline(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(7.4, 5.0), dpi=260)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8)
    ax.axis("off")
    draw_box(ax, 0.3, 5.8, 2.8, 1.0, "Client review", "Optional written comment")
    draw_box(ax, 4.6, 5.8, 2.8, 1.0, "Protected Edge Function", "Ownership check and job claim")
    draw_box(ax, 8.9, 5.8, 2.8, 1.0, "Python analysis API", "Shared-key protected endpoint")
    arrow(ax, (3.1, 6.3), (4.6, 6.3))
    arrow(ax, (7.4, 6.3), (8.9, 6.3))
    draw_box(ax, 0.3, 3.4, 2.8, 1.0, "Normalize comment", "Preserve language and sentiment cues")
    draw_box(ax, 4.6, 3.4, 2.8, 1.0, "Infer topic and retrieve", "LDA + six FAISS references")
    draw_box(ax, 8.9, 3.4, 2.8, 1.0, "Structured verdict", "Pre-trained OpenAI provider")
    ax.plot([10.3, 10.3, 1.7], [5.8, 5.1, 5.1], color="black", linewidth=1.1)
    arrow(ax, (1.7, 5.1), (1.7, 4.4))
    arrow(ax, (3.1, 3.9), (4.6, 3.9))
    arrow(ax, (7.4, 3.9), (8.9, 3.9))
    draw_box(ax, 4.6, 1.0, 2.8, 1.0, "Persist result", "Label · signed score · topic · provenance")
    ax.plot([10.3, 10.3, 6], [3.4, 2.6, 2.6], color="black", linewidth=1.1)
    arrow(ax, (6, 2.6), (6, 2.0))
    draw_box(ax, 0.3, 1.0, 2.8, 1.0, "Service page", "Ratings and sentiment groups")
    arrow(ax, (4.6, 1.5), (3.1, 1.5))
    fig.tight_layout(pad=0.4)
    fig.savefig(path, dpi=260, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.start_type = WD_SECTION_START.NEW_PAGE

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.italic = False

    if "Equation Explanation" not in [style.name for style in styles]:
        equation_style = styles.add_style("Equation Explanation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation_style = styles["Equation Explanation"]
    equation_style.font.name = "Times New Roman"
    equation_style.font.size = Pt(10)
    equation_style.font.italic = True
    equation_style.paragraph_format.left_indent = Inches(0.3)
    equation_style.paragraph_format.right_indent = Inches(0.3)
    equation_style.paragraph_format.space_after = Pt(7)

    header = section.header.paragraphs[0]
    header.text = "MULTIVENT — 3.2 Machine Learning Methodology"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_figure(document: Document, image_path: Path, caption: str, width=6.6) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    add_caption(document, caption, above=False)


def build_document() -> Document:
    create_training_pipeline(PIPELINE_IMAGE)
    create_runtime_pipeline(RUNTIME_IMAGE)

    document = Document()
    configure_document(document)
    document.core_properties.title = "3.2 Machine Learning Methodology — MULTIVENT Sentiment Model"
    document.core_properties.subject = "Dataset, preprocessing, algorithms, training, testing, and evaluation"
    document.core_properties.author = "MULTIVENT Capstone Team"
    document.core_properties.keywords = "MULTIVENT, sentiment analysis, LDA, RAG, FAISS, NLP"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("3.2 Machine Learning Methodology")

    add_body(
        document,
        "MULTIVENT applies a hybrid natural language processing pipeline to optional written "
        "comments submitted by clients after a booked service has been completed. The objective is "
        "to classify each non-empty comment as positive or negative, identify its dominant topic, "
        "retain traceable retrieval evidence, and present grouped feedback on the service page. Star "
        "ratings are stored and reported separately; they are not used to force the text sentiment label.",
    )
    add_body(
        document,
        "The term training is used carefully in this methodology. MULTIVENT fits an unsupervised "
        "Latent Dirichlet Allocation (LDA) topic model and constructs a training-only FAISS retrieval "
        "index. It does not train or fine-tune a supervised sentiment classifier. The recorded final "
        "sentiment verdict is produced by a configured pre-trained OpenAI model, augmented with LDA "
        "topic context and labeled examples retrieved from the local corpus. A weighted-neighbor "
        "classifier is implemented as a provisional local fallback.",
    )

    document.add_heading("3.2.1 Machine Learning Methodology", level=1)
    add_body(
        document,
        "The methodology separates offline corpus preparation from online inference. Offline processing "
        "registers and validates four documented datasets, standardizes binary labels, removes invalid "
        "and duplicate units, assigns leakage-safe partitions, and measures event-service relevance. "
        "Only domain-relevant training rows are used to fit LDA and create the balanced retrieval index. "
        "Validation and test records are never inserted into either training artifact.",
    )
    add_figure(document, PIPELINE_IMAGE, "Figure 3.2-1. MULTIVENT sentiment corpus preparation, model fitting, and evaluation pipeline.")
    add_body(
        document,
        "At runtime, the authenticated Supabase Edge Function verifies that the signed-in client owns "
        "the submitted review and then calls the protected Python analysis API. The API normalizes the "
        "comment, infers its LDA topic, retrieves six labeled examples, and requests a schema-constrained "
        "binary verdict. The resulting label, signed score, topic, provider, confidence, language, reason, "
        "and reference provenance are stored with the review. Retrieved training sentences themselves "
        "are not copied into the application database.",
    )
    add_figure(document, RUNTIME_IMAGE, "Figure 3.2-2. Runtime sentiment-analysis and application-integration flow.")

    document.add_heading("3.2.2 Dataset", level=1)
    add_body(
        document,
        "The corpus combines linguistic coverage from HiliSenti with service-domain coverage from "
        "restaurant and hospitality datasets. The registry pins repository revisions when available and "
        "records authorship, source URL, mapping rules, and license. No live marketplace website is "
        "scraped by the preparation pipeline.",
    )
    add_table(
        document,
        "Table 3.2-1. Registered source datasets.",
        ["Source", "Records", "Language/domain role", "License"],
        [
            ["HiliSenti v1", "23,337", "Hiligaynon, Filipino/Tagalog, English, and code-switched coverage", "CC BY-NC-SA 4.0"],
            ["MAMS ACSA", "3,000", "English restaurant and service aspects", "Apache 2.0"],
            ["Crawl Feeds Booking Hotel Reviews", "4,367", "English hospitality and venue/service reviews", "CC BY-NC 4.0"],
            ["UCI Sentiment Labelled Sentences — restaurant subset", "1,000", "English restaurant/service sentences", "CC BY 4.0"],
            ["Total", "31,704", "Original source records examined", "—"],
        ],
        widths=[1.8, 0.75, 3.05, 1.1],
        font_size=8.5,
    )
    add_body(
        document,
        "Source-record counts are not identical to candidate-text counts. One hotel review can yield "
        "separate positive and negative units. After schema validation, label mapping, exclusions, and "
        "exact deduplication, the complete binary corpus contains 23,171 valid text units. Neutral "
        "HiliSenti records (7,154) are excluded rather than relabeled.",
    )
    add_table(
        document,
        "Table 3.2-2. Generated corpus artifacts and their roles.",
        ["Artifact", "Rows", "Purpose"],
        [
            ["all_binary.csv", "23,171", "All valid, deduplicated positive/negative units"],
            ["domain_filtered.csv", "5,649", "Units meeting the event/service relevance threshold"],
            ["lda_train.csv", "4,462", "Training-only text used to fit LDA"],
            ["rag.csv", "3,050", "Balanced training-only references indexed by FAISS"],
            ["validation.csv", "2,376", "Held-out validation records"],
            ["evaluation_hilisenti.csv", "1,594", "Held-out HiliSenti test records"],
            ["evaluation_domain.csv", "790", "Held-out hospitality/restaurant test records"],
        ],
        widths=[2.0, 0.8, 3.9],
    )
    add_body(
        document,
        "The standardized CSV files contain 25 attributes: id, text_original, text_clean, sentiment, "
        "source_dataset, source_record_id, source_domain, language, split, aspect, license, source_url, "
        "text_hash, original_text_hash, source_name, language_basis, source_split, pool, label_basis, "
        "group_id, domain_relevance_score, domain_semantic_similarity, domain_keyword_score, "
        "domain_irrelevant_penalty, and domain_keyword_hits. These fields preserve the original text, "
        "target label, partition, source provenance, grouping information, and explainable domain-filtering evidence.",
    )
    add_body(
        document,
        "The main generated datasets are UTF-8 CSV files. Raw inputs include CSV, Parquet, and a ZIP "
        "archive containing the UCI restaurant text file. JSON stores manifests, checksums, parameters, "
        "and evaluation reports. LDA artifacts are serialized with joblib, while semantic references are "
        "stored in a FAISS binary index with JSON metadata.",
    )

    document.add_heading("3.2.3 Data Preprocessing", level=1)
    add_body(
        document,
        "Preprocessing is conservative because the corpus contains Hiligaynon, Filipino/Tagalog, "
        "English, and code-switched feedback. The pipeline avoids transformations that could remove "
        "negators, particles, emojis, punctuation, or informal forms that carry sentiment. Every "
        "exclusion and normalization step is recorded for auditability.",
    )
    add_bullets(document, [
        "Schema and label validation. Only supported labels are mapped to positive or negative. Neutral, conflicting, malformed, empty, placeholder, and explicit no-complaint units are excluded with reasons.",
        "Missing values. Required standardized fields are populated; invalid source units are excluded instead of imputed. A blank domain_keyword_hits value means that no configured keyword matched and is not a missing target.",
        "Text normalization. Unicode normalization and whitespace cleanup are applied. Reviewed token-boundary dictionaries expand selected informal Filipino and Hiligaynon shortcuts. Excessive repeated letters, limited laughter forms, and reduplication patterns are normalized with an audit trail.",
        "Cleaning and preservation. The provider/embedding text retains case, punctuation, emojis, particles, and the original language. PII masking for email addresses and Philippine mobile numbers exists but was disabled for the recorded corpus. Names and locations are not automatically anonymized.",
        "Duplicate control. Case-insensitive exact fingerprints are generated using SHA-256. Exact duplicates are removed using test-before-validation-before-training priority, and contradictory duplicate labels are quarantined.",
        "Feature encoding. LDA uses CountVectorizer bag-of-words counts. Semantic retrieval uses normalized 384-dimensional sentence-transformer vectors.",
        "Outlier handling. Numerical z-score or interquartile-range removal is not applicable to the text labels. Instead, semantically irrelevant text is filtered by an explainable relevance score; malformed and contradictory records are quarantined rather than numerically clipped.",
        "Leakage protection. Whole review groups remain in one partition. Protected test fingerprints are rejected from LDA and RAG, and index construction fails if an evaluation collision remains.",
    ])
    add_body(
        document,
        "For each comment vector v, L2 normalization is performed before similarity search. The small "
        "constant ε prevents division by zero.",
    )
    add_equation(document, "v̂ = v / max(‖v‖₂, ε)", 1,
                 "where v is the original embedding, v̂ is the normalized embedding, and ε is a small positive constant.")
    add_body(
        document,
        "Domain filtering combines the highest semantic similarity to reviewed domain descriptions, "
        "keyword evidence, a source-category prior, and an irrelevant-topic penalty.",
    )
    add_equation(document, "Kᵢ = min(1, hᵢ / 2)", 2,
                 "where hᵢ is the number of configured event/service keywords matched by text i.")
    add_equation(document, "Rᵢ = clip₍₀,₁₎(0.55Sᵢ + 0.30Kᵢ + 0.15Cᵢ − Pᵢ)", 3,
                 "Sᵢ is frozen-encoder semantic similarity, Kᵢ is keyword evidence, Cᵢ is the source-category prior, and Pᵢ is 0.30 when an irrelevant keyword is present. Training rows are retained when Rᵢ ≥ 0.40.")

    document.add_heading("3.2.4 Machine Learning Algorithms Used", level=1)
    document.add_heading("Latent Dirichlet Allocation", level=2)
    add_body(
        document,
        "LDA performs unsupervised topic discovery independently of sentiment labels. CountVectorizer "
        "lowercases internally, extracts word-character unigrams, and removes only 14 conservative "
        "English function words. Filipino and Hiligaynon stop-word lists remain empty so negators and "
        "sentiment-bearing particles are retained. No stemming or lemmatization is implemented.",
    )
    add_equation(document, "θ_d ~ Dirichlet(α),    φ_k ~ Dirichlet(β)", 4,
                 "θ_d represents the topic distribution of document d; φ_k represents the word distribution of topic k.")
    add_equation(document, "z_dn ~ Categorical(θ_d),    w_dn ~ Categorical(φ_zdn)", 5,
                 "For token n in document d, z_dn selects a topic and w_dn is generated from that topic's word distribution.")
    add_equation(document, "k* = arg max_k P(topic = k | document d)", 6,
                 "The runtime topic is the topic with the largest inferred posterior probability.")
    add_table(
        document,
        "Table 3.2-3. Recorded LDA configuration.",
        ["Parameter", "Value"],
        [
            ["Training rows", "4,462 training-only, domain-relevant units"],
            ["Number of topics", "10"],
            ["Maximum vocabulary", "10,000 features"],
            ["Document-frequency bounds", "min_df = 2; max_df = 0.95"],
            ["N-gram range", "Unigrams"],
            ["Learning method", "Batch"],
            ["Maximum iterations", "20"],
            ["Random state", "42"],
        ],
        widths=[2.4, 4.3],
    )
    add_body(
        document,
        "The generated topics remain labeled Unassigned Topic 0 through Unassigned Topic 9 because "
        "human validation of topic names has not yet been completed. Topic IDs and keywords are used "
        "as contextual evidence only and do not determine the sentiment label.",
    )

    document.add_heading("Multilingual Embeddings and FAISS Retrieval", level=2)
    add_body(
        document,
        "The frozen sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2 encoder produces "
        "384-dimensional vectors. The resolved model revision recorded in the evaluation artifact is "
        "e8f8c211226b894fcb81acc59f3b34ba3efd5f42. Long comments are divided into token windows, "
        "embedded, mean-pooled, and normalized. The model is used for inference only and is not fine-tuned.",
    )
    add_equation(document, "sim(q, xᵢ) = q̂ᵀx̂ᵢ = cos(q, xᵢ)", 7,
                 "Because query q and reference xᵢ are L2-normalized, FAISS inner product is equivalent to cosine similarity.")
    add_equation(document, "Wᵢ = max(0, sim(q, xᵢ)) × Rᵢ × ω_source", 8,
                 "Wᵢ is the reranking score, Rᵢ is domain relevance, and ω_source is 1.0 for HiliSenti or 1.2 for domain sources.")
    add_body(
        document,
        "The retrieval index contains 3,050 balanced training references: 1,525 positive and 1,525 "
        "negative. The default request retrieves six examples—two from the language pool and four from "
        "the domain pool—then sorts them by weighted evidence and cosine similarity.",
    )

    document.add_heading("Sentiment Decision", level=2)
    add_body(
        document,
        "The recorded active provider is a pre-trained OpenAI Responses model configured as "
        "gpt-5.6-terra. It receives the original feedback, inferred topic context, and retrieved examples. "
        "A Pydantic schema restricts the response to one binary label, positive and negative scores "
        "summing to one, confidence, language, and a concise reason. The prompt explicitly treats LDA "
        "and retrieved examples as context rather than ground truth.",
    )
    add_body(
        document,
        "The local fallback computes class evidence from the same reranked references. This fallback is "
        "implemented for resilience and research comparison; its evidence shares are not calibrated probabilities.",
    )
    add_equation(document, "E_c = Σ_{i:yᵢ=c} Wᵢ,    p_c = E_c / (E_positive + E_negative)", 9,
                 "The fallback selects positive only when p_positive > p_negative; a tie is resolved as negative.")
    add_equation(document, "sentiment_score = p_positive − p_negative,    −1 ≤ sentiment_score ≤ 1", 10,
                 "The signed score stored by MULTIVENT is positive evidence minus negative evidence. Provider confidence is retained separately and is not presented as a calibrated probability.")

    document.add_heading("3.2.5 Model Training and Testing", level=1)
    add_body(
        document,
        "HiliSenti retains its publisher-provided train, validation, and test partitions. MAMS, Booking "
        "Hotel Reviews, and UCI records are assigned by deterministic SHA-256 hashing of the whole "
        "review group with seed 42. Grouping prevents the positive and negative fields from one hotel "
        "review from entering different partitions.",
    )
    add_equation(document, "u_g = int(SHA256(\"42:\" ∥ group_id_g)[0:16], 16) / 2⁶⁴", 11,
                 "For non-HiliSenti sources: u_g < 0.10 gives test; 0.10 ≤ u_g < 0.20 gives validation; otherwise the group is training data.")
    add_table(
        document,
        "Table 3.2-4. Realized split of the complete binary corpus.",
        ["Partition", "Rows", "Percentage", "Use"],
        [
            ["Training", "18,411", "79.46%", "Source pool from which domain-relevant LDA/RAG subsets are drawn"],
            ["Validation", "2,376", "10.25%", "Held-out development validation"],
            ["Testing", "2,384", "10.29%", "Final held-out HiliSenti and domain-proxy evaluation pools"],
            ["Total", "23,171", "100.00%", "Valid deduplicated binary corpus"],
        ],
        widths=[1.15, 0.8, 0.85, 3.9],
    )
    add_body(
        document,
        "LDA is fitted on the 4,462 domain-relevant training rows without using their sentiment labels. "
        "FAISS is built from 3,050 training-only records balanced by class within each source using seed "
        "42. The index includes 1,525 positive and 1,525 negative references. Test and validation "
        "fingerprints are protected, and the generated manifest records that leakage checks passed.",
    )
    add_body(
        document,
        "No k-fold or stratified cross-validation is implemented. Fixed partitions were selected because "
        "the experiment preserves publisher splits, groups fields from the same review, builds retrieval "
        "artifacts from training records, and enforces strict fingerprint isolation. Therefore, the study "
        "must not claim cross-validation. Reproducibility is instead supported by deterministic hashing, "
        "seed 42, artifact checksums, and recorded evaluation samples.",
    )
    add_body(
        document,
        "The recorded evaluation used reproducible seed-42 samples of 100 held-out records from each of "
        "the two evaluation pools. Each request used six retrieved references. The full test pools remain "
        "available—1,594 HiliSenti records and 790 hospitality/restaurant records—but the recorded report "
        "did not evaluate every row.",
    )

    document.add_heading("3.2.6 Performance Evaluation Metrics", level=1)
    add_body(
        document,
        "For the following equations, TP is a positive comment correctly predicted as positive, TN is a "
        "negative comment correctly predicted as negative, FP is a negative comment predicted as positive, "
        "and FN is a positive comment predicted as negative.",
    )
    add_equation(document, "Accuracy = (TP + TN) / (TP + TN + FP + FN)", 12)
    add_equation(document, "Precision = TP / (TP + FP)    and    Recall = TP / (TP + FN)", 13)
    add_equation(document, "F1 = 2 × (Precision × Recall) / (Precision + Recall)", 14)
    add_equation(document, "Macro F1 = (F1_positive + F1_negative) / 2", 15,
                 "Macro F1 gives equal importance to both sentiment classes even when their sample counts differ.")
    add_table(
        document,
        "Table 3.2-5. Recorded binary sentiment evaluation.",
        ["Held-out set", "N", "Accuracy", "Precision", "Recall", "F1", "Macro F1"],
        [
            ["HiliSenti", "100", "0.9500", "0.9130", "0.9767", "0.9438", "0.9494"],
            ["Hospitality/restaurant domain proxy", "100", "0.9400", "0.9583", "0.9200", "0.9388", "0.9400"],
        ],
        widths=[2.1, 0.45, 0.8, 0.8, 0.7, 0.7, 0.85],
        font_size=8.5,
    )
    add_table(
        document,
        "Table 3.2-6. Confusion matrices from the recorded evaluation.",
        ["Set", "TN", "FP", "FN", "TP", "Correct / 100"],
        [
            ["HiliSenti", "53", "4", "1", "42", "95"],
            ["Domain proxy", "48", "2", "4", "46", "94"],
        ],
        widths=[2.0, 0.65, 0.65, 0.65, 0.65, 1.5],
    )
    add_body(
        document,
        "For HiliSenti, precision is 42/(42+4)=0.9130 and recall is 42/(42+1)=0.9767. "
        "For the domain proxy, precision is 46/(46+2)=0.9583 and recall is 46/(46+4)=0.9200. "
        "The reported macro F1 values of 0.9494 and 0.9400 indicate similar aggregate performance "
        "across both sentiment classes within these samples.",
    )

    document.add_heading("Interpretation and Limitations", level=2)
    add_body(
        document,
        "The recorded results are promising evidence that the configured hybrid pipeline can classify "
        "the sampled held-out reviews. They do not establish final performance on real MULTIVENT event "
        "feedback. Only 100 examples per pool were evaluated; the domain pool is a hospitality and "
        "restaurant proxy; some hotel labels are weak labels derived from positive/negative source fields; "
        "and the final provider was not fine-tuned on the assembled corpus.",
    )
    add_body(
        document,
        "Before sentiment is used for provider ranking, penalties, or automated business decisions, the "
        "study should collect a manually labeled and representative MULTIVENT validation set, establish "
        "annotation agreement, evaluate the complete held-out sets or justify sampling statistically, "
        "report per-language and per-service-category errors, and define acceptance thresholds before "
        "examining the final benchmark. Hiligaynon normalization and LDA topic names should also be "
        "reviewed by qualified speakers and domain experts.",
    )

    document.add_heading("Software and Implementation Tools", level=2)
    add_table(
        document,
        "Table 3.2-7. Tools used by the implemented ML/NLP pipeline.",
        ["Layer", "Implemented tools"],
        [
            ["Data preparation", "Python, pandas, NumPy, PyYAML, Hugging Face/UCI source downloads"],
            ["Topic model", "scikit-learn CountVectorizer and LatentDirichletAllocation; joblib artifacts"],
            ["Semantic retrieval", "Sentence Transformers, PyTorch-backed encoder, FAISS CPU"],
            ["Decision and API", "OpenAI Responses API, Pydantic structured output, FastAPI, Uvicorn"],
            ["Evaluation", "scikit-learn metrics, JSON/CSV reports, pytest and httpx"],
            ["System integration", "Supabase PostgreSQL, RPCs, Edge Functions and secrets; React Native/Expo"],
        ],
        widths=[1.55, 5.15],
    )

    document.add_heading("Evidence Used for This Methodology", level=2)
    add_body(
        document,
        "This section was prepared from the implemented training workspace and application integration. "
        "The principal evidence consisted of config/datasets.yaml; the generated corpus manifest; the "
        "dataset and binary evaluation reports; the LDA, retrieval, embedding, domain-scoring, and provider "
        "source files; and the MULTIVENT Supabase sentiment integration. The recorded evidence date for "
        "the evaluation artifacts is September 28, 2026.",
    )
    add_bullets(document, [
        "Jarder, J. J. T. HiliSenti v1. DOI: 10.57967/hf/8737.",
        "Jiang, Q., Chen, L., Xu, R., Ao, X., and Yang, M. (2019). A Challenge Dataset and Effective Models for Aspect-Based Sentiment Analysis.",
        "Kotzias, D. (2015). Sentiment Labelled Sentences. UCI Machine Learning Repository. DOI: 10.24432/C57604.",
        "Blei, D. M., Ng, A. Y., and Jordan, M. I. (2003). Latent Dirichlet Allocation. Journal of Machine Learning Research, 3, 993–1022.",
        "Reimers, N., and Gurevych, I. (2019). Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks.",
        "Johnson, J., Douze, M., and Jégou, H. (2017). Billion-scale similarity search with GPUs.",
    ])

    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
