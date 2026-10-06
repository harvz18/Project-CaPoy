"""Generate the evidence-based Chapter 3.2 machine-learning methodology Q&A."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "3.2_Machine_Learning_Methodology.docx"

WINE = "7C2438"
GOLD = "B88A44"
LIGHT_WINE = "F6EAED"
LIGHT_GRAY = "F2F3F5"
TEXT = RGBColor(39, 43, 48)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], WINE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            set_cell_text(cells[column_index], value)
            if row_index % 2:
                shade(cells[column_index], LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def add_answer(document: Document, number: int, question: str, paragraphs: list[str],
               bullets: list[str] | None = None) -> None:
    document.add_heading(f"{number}. {question}", level=2)
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(6)
    for item in bullets or []:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(27)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(WINE)

    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.size = Pt(13)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(GOLD)

    for name, size in (("Heading 1", 18), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(WINE)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "MULTIVENT  |  3.2 Machine Learning Methodology"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(WINE)
    add_page_number(section.footer.paragraphs[0])


def build_document() -> Document:
    document = Document()
    configure_document(document)

    document.core_properties.title = "3.2 Machine Learning Methodology"
    document.core_properties.subject = "MULTIVENT sentiment analysis, LDA, and RAG methodology"
    document.core_properties.author = "MULTIVENT Capstone Team"
    document.core_properties.keywords = "MULTIVENT, NLP, sentiment analysis, LDA, RAG, FAISS"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("3.2 Machine Learning Methodology")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("MULTIVENT Event-Services Marketplace")

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Evidence reviewed: October 4, 2026\n").bold = True
    meta.add_run("Model workspace: D:\\Sentiment Training\\multivent-sentiment\n")
    meta.add_run("Application workspace: D:\\Multivent\\Project-CaPoy\\Capstone")

    document.add_paragraph()
    note_table = document.add_table(rows=1, cols=1)
    note_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    note_table.style = "Table Grid"
    note_cell = note_table.cell(0, 0)
    shade(note_cell, LIGHT_WINE)
    note_cell.text = ""
    note_heading = note_cell.paragraphs[0].add_run("Implementation status\n")
    note_heading.bold = True
    note_heading.font.color.rgb = RGBColor.from_string(WINE)
    note_cell.paragraphs[0].add_run(
        "MULTIVENT fits an LDA topic model and builds a training-only FAISS retrieval index. "
        "It does not train or fine-tune a supervised sentiment classifier. The current final "
        "positive/negative verdict is produced by a configured pre-trained OpenAI model using "
        "the original comment, LDA topic context, and retrieved labeled examples."
    )

    document.add_paragraph()
    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.add_run(
        "This document answers the 22 methodology questions from the current source code, "
        "generated manifests, and recorded evaluation reports."
    ).italic = True
    document.add_page_break()

    document.add_heading("Methodology Questions and Answers", level=1)

    add_answer(
        document, 1, "What machine learning/NLP features are used in MULTIVENT?",
        ["MULTIVENT uses a hybrid NLP pipeline for optional written service-review comments."],
        [
            "Conservative Unicode, whitespace, and informal Filipino/Hiligaynon normalization.",
            "Bag-of-words count features for unsupervised LDA topic discovery.",
            "Pre-trained multilingual sentence embeddings with 384 dimensions.",
            "FAISS cosine-similarity retrieval of labeled training examples (RAG).",
            "Structured binary sentiment classification: positive or negative.",
            "Grounded positive/negative review-group summarization.",
            "Auditable topic, reference, score, provider, language, and reason metadata.",
        ],
    )

    add_answer(
        document, 2, "What is the purpose of the machine learning model in your system?",
        [
            "Its purpose is to turn a client's optional service-review comment into useful, "
            "traceable feedback. The pipeline assigns a positive or negative label, identifies "
            "the dominant topic, retrieves comparable labeled examples, and supports a short "
            "grounded summary for the service page.",
            "The result helps clients and providers understand recurring strengths and areas for "
            "improvement. Star ratings remain separate and are not used to force the text label. "
            "The output is informational and is not currently approved for automatic ranking, "
            "penalties, or other high-impact business decisions.",
        ],
    )

    add_answer(
        document, 3, "What dataset will you use?",
        [
            "The current corpus combines four registered sources. HiliSenti supplies Hiligaynon "
            "and code-switched coverage, while MAMS, Booking Hotel Reviews, and the UCI restaurant "
            "subset supply service, restaurant, and hospitality examples. Neutral records are "
            "excluded from this binary positive/negative experiment rather than relabeled."
        ],
    )
    add_table(
        document,
        ["Dataset", "Original records", "Primary role"],
        [
            ["HiliSenti v1", "23,337", "Hiligaynon/code-switched coverage"],
            ["MAMS ACSA", "3,000", "Restaurant/service coverage"],
            ["Crawl Feeds Booking Hotel Reviews", "4,367", "Hospitality coverage"],
            ["UCI restaurant sentiment subset", "1,000", "Restaurant/service coverage"],
        ],
        widths=[2.6, 1.1, 2.8],
    )

    add_answer(
        document, 4, "Where did/will you get the dataset?",
        [
            "HiliSenti v1 and the mirrored MAMS ACSA dataset were downloaded from Hugging Face at "
            "pinned revisions. The published Booking Hotel Reviews sample was also obtained from "
            "Hugging Face. The restaurant subset came from the UCI Machine Learning Repository's "
            "Sentiment Labelled Sentences archive. The project records each source URL, author, "
            "citation, revision when available, and license in config/datasets.yaml and the generated "
            "dataset-sources report. No live marketplace scraping is performed by this pipeline."
        ],
        [
            "HiliSenti: https://huggingface.co/datasets/jjjardev/hilisenti-v1 (DOI 10.57967/hf/8737).",
            "MAMS ACSA mirror: https://huggingface.co/datasets/justinsiow/mams-acsa; original project by Jiang et al.",
            "Booking Hotel Reviews: https://huggingface.co/datasets/crawlfeeds/Booking-Hotel-Reviews-Dataset.",
            "UCI source: https://archive.ics.uci.edu/dataset/331/sentiment+labelled+sentences (DOI 10.24432/C57604).",
        ],
    )

    add_answer(
        document, 5, "How many records are in the dataset?",
        [
            "The pipeline examined 31,704 original source records. After label mapping, exclusions, "
            "and exact deduplication, all_binary.csv contains 23,171 valid binary text units. The "
            "difference between source records and text units matters because one hotel review can "
            "supply separate positive and negative fields.",
            "For actual modeling, 5,649 units passed the event/service domain filter; 4,462 "
            "training-only rows fit LDA; and 3,050 balanced training-only rows form the FAISS/RAG "
            "index. There are 2,376 validation rows and 2,384 held-out test rows (1,594 HiliSenti "
            "plus 790 domain rows).",
        ],
    )
    add_table(
        document,
        ["Generated artifact", "Rows", "Purpose"],
        [
            ["all_binary.csv", "23,171", "Valid, deduplicated binary units"],
            ["domain_filtered.csv", "5,649", "Event/service-relevant units"],
            ["lda_train.csv", "4,462", "LDA fitting"],
            ["rag.csv", "3,050", "Balanced retrieval index"],
            ["validation.csv", "2,376", "Held-out validation"],
            ["evaluation_hilisenti.csv", "1,594", "HiliSenti test"],
            ["evaluation_domain.csv", "790", "Domain-proxy test"],
        ],
    )

    add_answer(
        document, 6, "What are the attributes/columns in the dataset?",
        [
            "The standardized processed files have 25 columns. They can be grouped as follows:"
        ],
        [
            "Identity and text: id, text_original, text_clean, text_hash, original_text_hash.",
            "Target and partition: sentiment, split, source_split.",
            "Source provenance: source_dataset, source_record_id, source_domain, source_name, source_url, license, group_id.",
            "Language and labeling: language, language_basis, aspect, pool, label_basis.",
            "Domain-scoring audit: domain_relevance_score, domain_semantic_similarity, domain_keyword_score, domain_irrelevant_penalty, domain_keyword_hits.",
        ],
    )

    add_answer(
        document, 7, "What file format is the dataset? (CSV, Excel, JSON, etc.)",
        [
            "The primary processed datasets used by LDA, RAG, validation, and evaluation are UTF-8 "
            "CSV files. Raw sources include CSV files, one Parquet file for MAMS, and a ZIP archive "
            "containing a tab-separated UCI text file. JSON files store manifests, checksums, source "
            "metadata, and evaluation results. Trained artifacts use joblib for LDA/vectorizer files "
            "and a FAISS binary index for retrieval."
        ],
    )

    add_answer(
        document, 8, "Does the dataset contain missing values?",
        [
            "Not in the required fields of the cleaned standardized corpus: identifiers, review "
            "text, binary sentiment, source, split, and required provenance are populated. Invalid "
            "or empty source units are excluded with audit reasons instead of being imputed. For "
            "example, 3,736 empty or placeholder hotel fields were excluded.",
            "One optional audit column, domain_keyword_hits, is blank when no configured keyword "
            "matched. It is blank in 16,213 of the 23,171 all-binary rows. This denotes no keyword "
            "hit and is not a missing label or missing review."
        ],
    )

    add_answer(
        document, 9, "What data preprocessing steps will you perform?",
        ["The implemented preparation process performs the following steps:"],
        [
            "Download and checksum pinned source files; retain source and license metadata.",
            "Validate schemas and map only supported labels to positive or negative.",
            "Exclude neutral, invalid, empty, placeholder, conflicting, and unsupported units with reasons.",
            "Apply Unicode normalization and whitespace cleanup.",
            "Apply reviewed token-boundary informal-language mappings and conservative reduplication rules.",
            "Create case-insensitive normalized fingerprints; remove exact duplicates and quarantine label conflicts.",
            "Assign or preserve train/validation/test groups before filtering and balancing.",
            "Compute domain relevance using frozen embeddings and reviewed keyword/category signals.",
            "Keep only relevant training rows for LDA and RAG; balance RAG classes within each source.",
            "Run fail-closed checks so held-out evaluation text cannot enter LDA or FAISS.",
        ],
    )

    add_answer(
        document, 10, "Will you remove duplicates, punctuation, stop words, emojis, or unnecessary characters?",
        [
            "Exact duplicate text is removed using a case-insensitive hash of Unicode-normalized, "
            "whitespace-normalized text. Contradictory duplicate labels are quarantined. Held-out "
            "records have priority so a duplicate cannot leak from test into training.",
            "Punctuation and emojis are not globally removed; the provider and embedding text preserve "
            "them because they can carry sentiment. Excess whitespace and a limited set of excessive "
            "repeated-letter/laughter patterns are normalized. For LDA only, 14 conservative English "
            "articles, prepositions, and conjunctions are removed. Filipino and Hiligaynon stop-word "
            "lists are intentionally empty so negators and sentiment-bearing particles are retained."
        ],
    )

    add_answer(
        document, 11, "Will you convert text to lowercase?",
        [
            "The stored original and clean text are not forced to lowercase. This preserves the input "
            "for the sentiment provider and embeddings. LDA's scikit-learn CountVectorizer lowercases "
            "internally when generating count features. Duplicate matching uses casefolding only for "
            "identity comparison; it does not overwrite the stored text."
        ],
    )

    add_answer(
        document, 12, "Will you use tokenization, stemming, or lemmatization?",
        [
            "Yes, tokenization is used. LDA uses CountVectorizer word-token extraction with the pattern "
            "\\b\\w+\\b and unigram features. The sentence-transformer uses its pre-trained subword "
            "tokenizer; long comments are divided into token windows and their normalized embeddings "
            "are mean-pooled.",
            "No stemming or lemmatization is implemented. This is deliberate because aggressive "
            "language-specific reduction could damage Hiligaynon, Filipino, English, and code-switched "
            "forms without validated multilingual rules."
        ],
    )

    add_answer(
        document, 13, "What machine learning algorithm(s) will you use?",
        ["The implemented pipeline uses several algorithms, each with a separate role:"],
        [
            "Latent Dirichlet Allocation (scikit-learn, batch learning) for unsupervised topic discovery.",
            "CountVectorizer bag-of-words counts as the LDA feature representation.",
            "A frozen paraphrase-multilingual-MiniLM-L12-v2 sentence-transformer for 384-dimensional semantic embeddings.",
            "FAISS inner-product search over normalized vectors for cosine-similarity retrieval.",
            "A pre-trained OpenAI Responses model for the current structured binary sentiment verdict and grounded summaries.",
            "A provisional weighted nearest-reference vote as the local fallback provider.",
        ],
    )

    add_answer(
        document, 14, "Are you using LDA (Latent Dirichlet Allocation)? If yes, how?",
        [
            "Yes. LDA is trained only on 4,462 domain-relevant training rows and uses text counts, not "
            "sentiment labels. The recorded configuration is 10 topics, up to 10,000 unigram features, "
            "min_df=2, max_df=0.95, batch learning, 20 maximum iterations, and random state 42.",
            "At inference, the vectorizer transforms the normalized comment, LDA produces a topic "
            "distribution, and the highest-probability topic plus its top keywords is returned. If no "
            "known vocabulary is present, the system reports no known topic instead of inventing one. "
            "Topic context is included in the final request but does not determine sentiment. Topic "
            "names remain unassigned pending human review."
        ],
    )

    add_answer(
        document, 15, "What sentiment analysis method/algorithm will you use?",
        [
            "The current production-intended method is contextual binary classification through a "
            "schema-constrained OpenAI Responses call. It receives the original comment, the LDA topic, "
            "and six retrieved training examples (two language-pool and four domain-pool examples in "
            "the current settings). It must return exactly positive or negative, two scores summing to "
            "one, confidence, detected/estimated language, and a concise reason.",
            "If the local fallback is selected, the system sums max(0, cosine similarity) x domain "
            "relevance x source weight for positive and negative retrieved references and chooses the "
            "larger total. The fallback share is explicitly not a calibrated probability."
        ],
    )

    add_answer(
        document, 16, "Is your sentiment analysis model trained using your dataset, or are you using a pre-trained model/library?",
        [
            "The final sentiment classifier is a pre-trained OpenAI model configured as gpt-5.6-terra; "
            "it is not trained or fine-tuned on the assembled MULTIVENT corpus. The sentence-transformer "
            "is also pre-trained and frozen.",
            "The local corpus is still essential: it is used to fit LDA and to build the balanced FAISS "
            "retrieval index that supplies relevant labeled examples. Therefore, the accurate description "
            "is a pre-trained sentiment provider augmented by locally fitted topic discovery and "
            "training-only retrieval, not a newly trained supervised sentiment model."
        ],
    )

    add_answer(
        document, 17, "How will you divide the dataset into training and testing data?",
        [
            "HiliSenti retains its publisher-provided train, validation, and test partitions. The other "
            "sources are divided by deterministic SHA-256 hashing of whole review groups with seed 42. "
            "This keeps multiple fields from the same hotel review together.",
            "Splitting occurs before domain filtering and class balancing. Exact duplicates prioritize "
            "test, then validation, then training; conflicting labels are quarantined. LDA and FAISS use "
            "training-only data, and automated guards fail the build if held-out fingerprints or source "
            "groups enter training."
        ],
    )

    add_answer(
        document, 18, "What percentage will be used for training and testing?",
        [
            "For MAMS, Booking Hotel Reviews, and UCI, the configured split is 80% training, 10% "
            "validation, and 10% testing by deterministic group assignment. HiliSenti keeps its original "
            "source partitions rather than being re-split.",
            "Across the current 23,171-row binary corpus after deduplication, the realized totals are "
            "18,411 training rows (79.46%), 2,376 validation rows (10.25%), and 2,384 test rows (10.29%). "
            "Only domain-relevant training subsets are used for LDA (4,462) and balanced RAG (3,050)."
        ],
    )

    add_answer(
        document, 19, "Will you use cross-validation? If yes, what type?",
        [
            "No k-fold or stratified cross-validation is currently implemented. The methodology uses "
            "fixed train, validation, and test partitions with whole-review grouping, source-preserved "
            "HiliSenti splits, reproducible seed-42 sampling, and explicit leakage guards. The held-out "
            "test set is reserved for final evaluation; development or model-selection experiments "
            "should use validation data without adding it to the RAG index."
        ],
    )

    add_answer(
        document, 20, "What performance evaluation metrics will you use? (Accuracy, Precision, Recall, F1-score, etc.)",
        [
            "The evaluation script reports accuracy, positive-class precision, positive-class recall, "
            "positive-class F1, macro precision, macro recall, macro F1, and a 2 x 2 confusion matrix "
            "with negative and positive labels. It also records metrics by source, class/source "
            "distribution, individual predictions, elapsed time, configuration, artifact hashes, and "
            "whether leakage checks passed. Macro F1 is important because it weights both classes "
            "equally even when their counts differ."
        ],
    )

    add_answer(
        document, 21, "What results will you use to determine whether the model performs well?",
        [
            "The current reproducible report used 100 held-out records from each evaluation pool. "
            "HiliSenti achieved 0.9500 accuracy and 0.9494 macro F1. The hospitality/restaurant proxy "
            "set achieved 0.9400 accuracy and 0.9400 macro F1. Leakage checks passed."
        ],
    )
    add_table(
        document,
        ["Held-out set", "N", "Accuracy", "Precision", "Recall", "F1", "Macro F1"],
        [
            ["HiliSenti", "100", "0.9500", "0.9130", "0.9767", "0.9438", "0.9494"],
            ["Domain proxy", "100", "0.9400", "0.9583", "0.9200", "0.9388", "0.9400"],
        ],
    )
    for text in [
        "These scores are promising but are not sufficient by themselves to claim strong performance "
        "on real MULTIVENT feedback. The samples are limited to 100 per set; the domain set is a "
        "hospitality/restaurant proxy; and some hotel labels use weak field-position supervision.",
        "The final determination should use the full untouched test sets, confusion matrices, macro F1, "
        "per-source/per-language results, qualitative error analysis, and a separately collected, "
        "manually labeled MULTIVENT event-service set. Acceptance thresholds must be defined before that "
        "final evaluation. Sentiment should not drive ranking or penalties until this validation is complete.",
    ]:
        document.add_paragraph(text)

    add_answer(
        document, 22, "What programming language, libraries, or tools will you use for the machine learning/NLP implementation?",
        ["The implementation uses the following languages, libraries, and services:"],
        [
            "Python for the ML/NLP service and preparation scripts.",
            "pandas and NumPy for data handling; scikit-learn for CountVectorizer, LDA, and metrics; joblib for artifacts.",
            "Sentence Transformers for embeddings and FAISS CPU for vector retrieval.",
            "FastAPI, Uvicorn, Pydantic, python-dotenv, and PyYAML for the API and configuration.",
            "OpenAI's Python SDK and Responses API for schema-constrained verdicts and summaries.",
            "Hugging Face Hub/Datasets and UCI downloads for source acquisition.",
            "pytest and httpx for automated validation.",
            "Supabase PostgreSQL, RPCs, Edge Functions, and secrets for secure application integration.",
            "TypeScript and React Native/Expo for the MULTIVENT mobile interface.",
        ],
    )

    document.add_page_break()
    document.add_heading("Evidence Base and Research Limitations", level=1)
    document.add_paragraph(
        "The answers above were derived from the current implementation and generated artifacts, not "
        "from a proposed architecture alone. The principal evidence files are listed below."
    )
    add_table(
        document,
        ["Evidence", "What it establishes"],
        [
            ["config/datasets.yaml", "Sources, revisions, licenses, mappings, and roles"],
            ["multivent_corpus/manifest.json", "Counts, partitions, settings, checksums, and leakage status"],
            ["reports/multivent_dataset_report.md", "Source exclusions and generated-corpus distributions"],
            ["reports/evaluation.json", "Provider, samples, metrics, matrices, hashes, and limitations"],
            ["app/corpus/prepare.py", "Splitting, deduplication, filtering, balancing, and leakage guards"],
            ["app/topics/lda_service.py", "LDA feature extraction, fitting, saving, and inference"],
            ["app/sentiment/openai_provider.py", "Structured sentiment and summary provider behavior"],
            ["docs/sentiment-integration.md", "Application-to-model runtime and deployment flow"],
        ],
    )
    document.add_paragraph(
        "Important limitation: the current report is not a final MULTIVENT field benchmark. A manually "
        "labeled production-representative dataset, native-speaker review, stable deployment, full test "
        "evaluation, and predefined acceptance criteria remain necessary."
    )

    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

