from __future__ import annotations

import math
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt

from generate_live_database_erd import (
    DOMAIN_SLUGS,
    DOMAIN_TABLES,
    EXPORT_DIR,
    PROJECT_REF,
    compact_type,
    domain_for_table,
    fk_column_map,
    load_schema,
    multiplicities,
    table_map,
)


STYLE_DIR = EXPORT_DIR / "database_style"
CANVAS_W = 4200
CANVAS_H = 2970
PDF_PATH = STYLE_DIR / "MULTIVENT_ERD_DATABASE_STYLE_MANUSCRIPT.pdf"
DOCX_PATH = STYLE_DIR / "MULTIVENT_ERD_DATABASE_STYLE_MANUSCRIPT.docx"

BG = (15, 15, 15)
GRID = (38, 38, 38)
CARD = (24, 24, 24)
CARD_ALT = (28, 28, 28)
HEADER = (31, 31, 31)
BORDER = (74, 74, 74)
LINE = (107, 107, 107)
TEXT = (238, 238, 238)
MUTED = (168, 168, 168)
KEY = (255, 255, 255)


def load_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


FONT_TITLE = load_font(56, True)
FONT_SUBTITLE = load_font(24)
FONT_SECTION = load_font(29, True)
FONT_HEADER = load_font(23, True)
FONT_COLUMN = load_font(18)
FONT_COLUMN_BOLD = load_font(18, True)
FONT_TYPE = load_font(15)
FONT_SMALL = load_font(15)
FONT_TINY = load_font(13)


def trim(draw: ImageDraw.ImageDraw, value: str, font, max_width: int) -> str:
    if draw.textlength(value, font=font) <= max_width:
        return value
    text = value
    while text and draw.textlength(text + "...", font=font) > max_width:
        text = text[:-1]
    return text.rstrip() + "..."


def draw_grid(draw: ImageDraw.ImageDraw) -> None:
    spacing = 34
    radius = 1
    for y in range(18, CANVAS_H, spacing):
        for x in range(18, CANVAS_W, spacing):
            draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=GRID)


def column_is_fk(table_name: str, column_name: str, fk_columns: dict[str, set[str]]) -> bool:
    return column_name in fk_columns.get(table_name, set())


def key_kind(table_name: str, column: dict, fk_columns: dict[str, set[str]]) -> str:
    is_fk = column_is_fk(table_name, column["name"], fk_columns)
    if column["primary_key"] and is_fk:
        return "PK/FK"
    if column["primary_key"]:
        return "PK"
    if is_fk:
        return "FK"
    if column["unique"]:
        return "UQ"
    return ""


def draw_column_icon(draw: ImageDraw.ImageDraw, x: int, y: int, kind: str) -> None:
    if kind in {"PK", "PK/FK"}:
        draw.ellipse((x - 5, y - 5, x + 5, y + 5), outline=KEY, width=2)
        draw.line((x + 5, y, x + 15, y), fill=KEY, width=2)
        draw.line((x + 11, y, x + 11, y + 5), fill=KEY, width=2)
        draw.line((x + 15, y, x + 15, y + 4), fill=KEY, width=2)
    elif kind in {"FK", "UQ"}:
        draw.polygon(
            [(x, y - 6), (x + 6, y), (x, y + 6), (x - 6, y)],
            outline=TEXT,
        )
    else:
        draw.ellipse((x - 4, y - 4, x + 4, y + 4), outline=MUTED, width=1)


def draw_table_icon(draw: ImageDraw.ImageDraw, x: int, y: int) -> None:
    draw.rectangle((x, y, x + 20, y + 20), outline=MUTED, width=2)
    draw.line((x, y + 7, x + 20, y + 7), fill=MUTED, width=1)
    draw.line((x + 7, y, x + 7, y + 20), fill=MUTED, width=1)


def layout_cards(names: list[str]) -> dict[str, tuple[int, int, int, int]]:
    cols = 4 if len(names) >= 7 else 3
    rows = math.ceil(len(names) / cols)
    margin_x = 95
    gap_x = 70
    top = 220
    bottom = 115
    gap_y = 65
    card_w = (CANVAS_W - 2 * margin_x - (cols - 1) * gap_x) // cols
    card_h = (CANVAS_H - top - bottom - (rows - 1) * gap_y) // rows
    result = {}
    for index, name in enumerate(names):
        row, col = divmod(index, cols)
        result[name] = (
            margin_x + col * (card_w + gap_x),
            top + row * (card_h + gap_y),
            card_w,
            card_h,
        )
    return result


def display_columns(table_name: str, table: dict, fk_columns, card_h: int) -> list[dict]:
    header_h = 52
    row_h = 29
    max_rows = max(4, (card_h - header_h - 14) // row_h)
    columns = table["columns"]
    if len(columns) <= max_rows:
        return columns

    required_names = {
        column["name"]
        for column in columns
        if column["primary_key"] or column_is_fk(table_name, column["name"], fk_columns)
    }
    selected_names = set(required_names)
    room_for_columns = max(max_rows - 1, len(required_names))
    for column in columns:
        if len(selected_names) >= room_for_columns:
            break
        if column["name"] not in selected_names:
            selected_names.add(column["name"])
    selected = [column for column in columns if column["name"] in selected_names]
    omitted = len(columns) - len(selected)
    if omitted > 0:
        selected.append(
            {
                "name": f"... +{omitted} additional columns",
                "type": "",
                "nullable": True,
                "primary_key": False,
                "unique": False,
                "_summary": True,
            }
        )
    return selected[:max_rows]


def build_cards(names, tables, fk_columns, positions):
    cards = {}
    for name in names:
        box = positions[name]
        columns = display_columns(name, tables[name], fk_columns, box[3])
        row_h = 29
        row_map = {
            column["name"]: box[1] + 52 + index * row_h + row_h // 2
            for index, column in enumerate(columns)
            if not column.get("_summary")
        }
        cards[name] = {
            "box": box,
            "columns": columns,
            "row_map": row_map,
            "row_h": row_h,
        }
    return cards


def relationship_route(source_card, target_card, source_column, target_column, edge_index):
    sx, sy, sw, sh = source_card["box"]
    tx, ty, tw, th = target_card["box"]
    source_y = source_card["row_map"].get(source_column, sy + 78)
    target_y = target_card["row_map"].get(target_column, ty + 78)
    if sx + sw < tx:
        start = (sx + sw, source_y)
        end = (tx, target_y)
        middle_x = (start[0] + end[0]) // 2 + ((edge_index % 7) - 3) * 5
        points = [start, (middle_x, source_y), (middle_x, target_y), end]
    elif tx + tw < sx:
        start = (sx, source_y)
        end = (tx + tw, target_y)
        middle_x = (start[0] + end[0]) // 2 + ((edge_index % 7) - 3) * 5
        points = [start, (middle_x, source_y), (middle_x, target_y), end]
    else:
        use_right = edge_index % 2 == 0
        if use_right:
            start = (sx + sw, source_y)
            end = (tx + tw, target_y)
            outside_x = max(sx + sw, tx + tw) + 30 + (edge_index % 8) * 9
        else:
            start = (sx, source_y)
            end = (tx, target_y)
            outside_x = min(sx, tx) - 30 - (edge_index % 8) * 9
        points = [start, (outside_x, source_y), (outside_x, target_y), end]
    return points


def draw_card(
    draw: ImageDraw.ImageDraw,
    table_name: str,
    table: dict,
    card: dict,
    fk_columns: dict[str, set[str]],
) -> None:
    x, y, width, height = card["box"]
    draw.rectangle((x, y, x + width, y + height), fill=CARD, outline=BORDER, width=2)
    draw.rectangle((x, y, x + width, y + 52), fill=HEADER, outline=BORDER, width=2)
    draw_table_icon(draw, x + 15, y + 16)
    draw.text(
        (x + 47, y + 26),
        trim(draw, table_name, FONT_HEADER, width - 85),
        anchor="lm",
        fill=TEXT,
        font=FONT_HEADER,
    )
    draw.ellipse((x + width - 27, y + 22, x + width - 21, y + 28), fill=MUTED)
    row_h = card["row_h"]
    for index, column in enumerate(card["columns"]):
        row_y = y + 52 + index * row_h
        fill = CARD_ALT if index % 2 else CARD
        draw.rectangle((x + 1, row_y, x + width - 1, row_y + row_h), fill=fill)
        draw.line((x, row_y + row_h, x + width, row_y + row_h), fill=(50, 50, 50), width=1)
        if column.get("_summary"):
            draw.text(
                (x + 18, row_y + row_h // 2),
                column["name"],
                anchor="lm", fill=MUTED, font=FONT_SMALL,
            )
            continue
        kind = key_kind(table_name, column, fk_columns)
        draw_column_icon(draw, x + 21, row_y + row_h // 2, kind)
        name_font = FONT_COLUMN_BOLD if kind else FONT_COLUMN
        nullable = " ?" if column["nullable"] else ""
        name_text = trim(
            draw,
            column["name"] + nullable,
            name_font,
            int(width * 0.61) - 55,
        )
        draw.text(
            (x + 43, row_y + row_h // 2),
            name_text,
            anchor="lm", fill=KEY if kind else TEXT, font=name_font,
        )
        type_text = compact_type(column["type"])
        draw.text(
            (x + width - 16, row_y + row_h // 2),
            trim(draw, type_text, FONT_TYPE, int(width * 0.34)),
            anchor="rm", fill=MUTED, font=FONT_TYPE,
        )


def draw_cardinality_chip(draw, point, toward, value):
    px, py = point
    tx, ty = toward
    if abs(tx - px) >= abs(ty - py):
        lx = px + (34 if tx > px else -34)
        ly = py - 13
    else:
        lx = px + 32
        ly = py + (26 if ty > py else -26)
    width = 50 if value != "0..*" else 60
    height = 24
    draw.rectangle(
        (lx - width // 2, ly - height // 2, lx + width // 2, ly + height // 2),
        fill=BG,
        outline=BORDER,
        width=1,
    )
    draw.text((lx, ly), value, anchor="mm", fill=TEXT, font=FONT_TINY)


def generate_section(schema, tables, fk_columns, domain, names):
    positions = layout_cards(names)
    cards = build_cards(names, tables, fk_columns, positions)
    internal_fks = [
        fk for fk in schema["foreign_keys"]
        if fk["source_table"] in cards and fk["target_table"] in cards
    ]
    image = Image.new("RGB", (CANVAS_W, CANVAS_H), BG)
    draw = ImageDraw.Draw(image)
    draw_grid(draw)
    draw.text(
        (CANVAS_W // 2, 54),
        f"MULTIVENT ERD - {domain}",
        anchor="ma", fill=TEXT, font=FONT_TITLE,
    )
    draw.text(
        (CANVAS_W // 2, 130),
        f"{len(names)} tables  |  {len(internal_fks)} relationships",
        anchor="ma", fill=MUTED, font=FONT_SUBTITLE,
    )

    routes = []
    for edge_index, fk in enumerate(internal_fks, start=1):
        points = relationship_route(
            cards[fk["source_table"]],
            cards[fk["target_table"]],
            fk["source_columns"][0],
            fk["target_columns"][0],
            edge_index,
        )
        routes.append((fk, points))
        draw.line(points, fill=LINE, width=2)

    for name in names:
        draw_card(draw, name, tables[name], cards[name], fk_columns)

    for fk, points in routes:
        draw.ellipse(
            (points[0][0] - 3, points[0][1] - 3, points[0][0] + 3, points[0][1] + 3),
            fill=TEXT,
        )
        draw.ellipse(
            (points[-1][0] - 3, points[-1][1] - 3, points[-1][0] + 3, points[-1][1] + 3),
            fill=TEXT,
        )
        child_mult, parent_mult = multiplicities(fk)
        draw_cardinality_chip(draw, points[0], points[1], child_mult)
        draw_cardinality_chip(draw, points[-1], points[-2], parent_mult)

    draw.text(
        (80, CANVAS_H - 58),
        "PK key icon  |  FK/UQ diamond  |  ? nullable  |  Cardinality: 1, 0..1, 0..*",
        anchor="lm", fill=MUTED, font=FONT_SMALL,
    )
    filename = f"MULTIVENT_ERD_DATABASE_STYLE_{DOMAIN_SLUGS[domain]}.png"
    path = STYLE_DIR / filename
    image.save(path, dpi=(300, 300))
    return path


def panel_route(source_box, target_box, index):
    sx, sy, sw, sh = source_box
    tx, ty, tw, th = target_box
    source = (sx + sw // 2, sy + sh // 2)
    target = (tx + tw // 2, ty + th // 2)
    if abs(target[0] - source[0]) > abs(target[1] - source[1]):
        if target[0] > source[0]:
            start = (sx + sw, source[1])
            end = (tx, target[1])
        else:
            start = (sx, source[1])
            end = (tx + tw, target[1])
        mid_x = (start[0] + end[0]) // 2 + ((index % 5) - 2) * 8
        return [start, (mid_x, start[1]), (mid_x, end[1]), end]
    if target[1] > source[1]:
        start = (source[0], sy + sh)
        end = (target[0], ty)
    else:
        start = (source[0], sy)
        end = (target[0], ty + th)
    mid_y = (start[1] + end[1]) // 2 + ((index % 5) - 2) * 8
    return [start, (start[0], mid_y), (end[0], mid_y), end]


def generate_overview(schema):
    image = Image.new("RGB", (CANVAS_W, CANVAS_H), BG)
    draw = ImageDraw.Draw(image)
    draw_grid(draw)
    draw.text(
        (CANVAS_W // 2, 54),
        "MULTIVENT ERD - Overview",
        anchor="ma", fill=TEXT, font=FONT_TITLE,
    )
    draw.text(
        (CANVAS_W // 2, 130),
        f"{len(schema['tables'])} tables  |  "
        f"{sum(len(t['columns']) for t in schema['tables'])} columns  |  "
        f"{len(schema['foreign_keys'])} foreign keys",
        anchor="ma", fill=MUTED, font=FONT_SUBTITLE,
    )
    positions = {
        "Identity and Access": (80, 260, 1040, 780),
        "Marketplace and Availability": (1240, 260, 1040, 900),
        "Booking and Finance": (2400, 260, 1040, 720),
        "Event Planning and Coordination": (490, 1510, 1400, 1000),
        "Communication, Feedback, and Operations": (2310, 1510, 1400, 1000),
    }
    counts = Counter()
    for fk in schema["foreign_keys"]:
        source_domain = domain_for_table(fk["source_table"])
        target_domain = domain_for_table(fk["target_table"])
        if source_domain != target_domain:
            counts[(source_domain, target_domain)] += 1
    for index, ((source_domain, target_domain), count) in enumerate(
        sorted(counts.items()), start=1
    ):
        points = panel_route(positions[source_domain], positions[target_domain], index)
        draw.line(points, fill=LINE, width=3)
        a, b = points[1], points[2]
        lx, ly = (a[0] + b[0]) // 2, (a[1] + b[1]) // 2
        draw.rectangle((lx - 43, ly - 18, lx + 43, ly + 18), fill=BG, outline=BORDER)
        draw.text((lx, ly), f"{count} FK", anchor="mm", fill=TEXT, font=FONT_TINY)
    for domain, names in DOMAIN_TABLES.items():
        x, y, width, height = positions[domain]
        draw.rectangle((x, y, x + width, y + height), fill=CARD, outline=BORDER, width=2)
        draw.rectangle((x, y, x + width, y + 58), fill=HEADER, outline=BORDER, width=2)
        draw_table_icon(draw, x + 18, y + 18)
        draw.text((x + 53, y + 29), domain, anchor="lm", fill=TEXT, font=FONT_SECTION)
        cols = 2
        gap = 18
        mini_w = (width - 44 - gap) // cols
        mini_h = 45
        start_y = y + 82
        for table_index, name in enumerate(names):
            row, col = divmod(table_index, cols)
            bx = x + 18 + col * (mini_w + gap)
            by = start_y + row * (mini_h + 12)
            draw.rectangle(
                (bx, by, bx + mini_w, by + mini_h),
                fill=HEADER,
                outline=(60, 60, 60),
                width=1,
            )
            draw_table_icon(draw, bx + 10, by + 12)
            draw.text(
                (bx + 40, by + mini_h // 2),
                trim(draw, name, FONT_SMALL, mini_w - 52),
                anchor="lm", fill=TEXT, font=FONT_SMALL,
            )
    path = STYLE_DIR / "MULTIVENT_ERD_DATABASE_STYLE_OVERVIEW.png"
    image.save(path, dpi=(300, 300))
    return path


def build_pdf(paths):
    images = [Image.open(path).convert("RGB") for path in paths]
    images[0].save(
        PDF_PATH,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=images[1:],
    )
    for image in images:
        image.close()


def build_docx(paths, captions):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(8)
    section.bottom_margin = Mm(8)
    section.left_margin = Mm(8)
    section.right_margin = Mm(8)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9)
    for index, (path, caption) in enumerate(zip(paths, captions)):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(caption)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(str(path), width=Inches(10.9))
        if index < len(paths) - 1:
            document.add_page_break()
    document.save(DOCX_PATH)


def main():
    STYLE_DIR.mkdir(parents=True, exist_ok=True)
    schema = load_schema()
    tables = table_map(schema)
    fk_columns = fk_column_map(schema)
    paths = [generate_overview(schema)]
    captions = ["MULTIVENT ERD - Overview"]
    for domain, names in DOMAIN_TABLES.items():
        paths.append(generate_section(schema, tables, fk_columns, domain, names))
        captions.append(f"MULTIVENT ERD - {domain}")
    build_pdf(paths)
    build_docx(paths, captions)
    print("Generated database-style ERD images:")
    for path in paths:
        print(path)
    print(PDF_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()



